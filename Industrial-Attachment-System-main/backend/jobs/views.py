import re
import csv
import math
import logging
import statistics
from datetime import date, datetime, timedelta
from django.db.models import Q, Count, Avg, Sum
from django.db.models.functions import Coalesce, TruncDate, TruncMonth, TruncWeek, TruncYear
from django.http import HttpResponse
from rest_framework import generics, permissions, status
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.throttling import UserRateThrottle
from .models import (
    Department,
    Requisition,
    Job,
    Application,
    ApplicationStatusEvent,
    Notification,
    Deployment,
    Clearance,
    RecommendationLetterTemplate,
    RecommendationLetter,
    EligibilitySetting,
    AtsScoringConfiguration,
)
from .attachment_records import (
    COMPLETED_ATTACHMENT_CONDITION,
    completed_attachment_queryset,
    is_completed_attachment,
    serialize_attachment_archive_row,
)
from .serializers import (
    DepartmentSerializer,
    RequisitionSerializer,
    JobSerializer,
    PublicJobSerializer,
    ApplicationSerializer,
    NotificationSerializer,
    DeploymentSerializer,
    ClearanceSerializer,
    RecommendationLetterTemplateSerializer,
    RecommendationLetterSerializer,
    EligibilitySettingSerializer,
    AtsScoringConfigurationSerializer,
)
from django.contrib.auth import get_user_model
from rest_framework.views import APIView
from rest_framework.response import Response
from django.utils import timezone
from django.db.models.functions import TruncDate, TruncMonth, TruncWeek, TruncYear

from .sanitizers import rich_text_to_plain_text

logger = logging.getLogger(__name__)

User = get_user_model()

# ---------------------------------------------------------------------------
# ATS (Applicant Tracking System) configuration
# ---------------------------------------------------------------------------

ENGLISH_STOP_WORDS = frozenset({
    "a", "an", "the", "and", "or", "but", "in", "on", "at", "to", "for",
    "of", "with", "by", "is", "are", "was", "were", "be", "been", "being",
    "have", "has", "had", "do", "does", "did", "will", "would", "could",
    "should", "may", "might", "shall", "this", "that", "these", "those",
    "it", "its", "from", "not", "no", "as", "we", "our", "your", "their",
    "which", "who", "you", "they", "he", "she", "me", "him", "her",
    "us", "them", "any", "all", "must", "can", "also", "up", "out", "into",
    "so", "if", "than", "then", "when", "where", "how", "what", "each",
    "both", "only", "just", "over", "per", "via", "s", "re", "ve", "ll",
    "d", "t", "m", "such", "other", "more", "very", "well", "some",
})

# The "minimum score" a candidate must reach is not a fixed constant: it is
# derived per-application from the Admin-configurable quality tier boundaries.
# See AtsScoringConfiguration.

DOCUMENT_ATS_KEYWORDS = {
    "NATIONAL_ID":       "national identity identification citizen id card kenya",
    "RESUME":            "resume cv curriculum vitae work experience skills profile professional",
    "COVER_LETTER":      "cover letter motivation statement application interest suitability",
    "INSTITUTION_INTRO": "introduction letter institution university recommendation attachment intern",
    "STUDENT_INSURANCE": "insurance student medical health cover protection policy",
    "STUDENT_ID":        "student identification university college enrollment registration card",
    "NEXT_OF_KIN_ID":    "next kin emergency contact family relative beneficiary",
    "TRANSCRIPT":        "transcript academic records marks grades results coursework examination",
    "GOOD_CONDUCT":      "conduct certificate police clearance character background vetting integrity",
    "PASSPORT_PHOTOS":   "passport photo identity biometric portrait",
    # "KRA_PIN" was removed from the required upload set (replaced by
    # TRANSCRIPT and GOOD_CONDUCT) and no longer contributes keywords.
    # Commented out per requirements:
    # "SHA_CARD":          "sha social health insurance authority card medical",
    # "NSSF_CARD":         "nssf national social security fund pension card contribution",
    # "BIRTH_CERT":        "birth certificate registration civil record",
    # "ACADEMIC_CERT":     "academic certificate qualification degree diploma award graduation",
    # "SECRETS_ACT_FORM":  "official secrets act security clearance government confidentiality",
    # "PSIP_FORM":         "psip intern biodata attachment public service internship form",
    # "ATM_CARD":          "atm bank account banking payment details financial",
}

# The ten mandatory verification documents. TRANSCRIPT and GOOD_CONDUCT replaced
# KRA_PIN; because every required document carries an equal share of
# ATS_DOCUMENT_WEIGHT, adding a required document lowers the per-document value
# and therefore shifts existing stored scores. See the `recompute_ats_scores`
# management command, which must be run after this set changes.
REQUIRED_DOCUMENTS = [
    "NATIONAL_ID",
    "RESUME",
    "COVER_LETTER",
    "INSTITUTION_INTRO",
    "STUDENT_INSURANCE",
    "STUDENT_ID",
    "NEXT_OF_KIN_ID",
    "TRANSCRIPT",
    "GOOD_CONDUCT",
    "PASSPORT_PHOTOS",
]
ATS_REQUIRED_DOC_COUNT = len(REQUIRED_DOCUMENTS)

# ATS scoring weights. These sum to exactly ATS_SCORE_MAX, which is the ceiling
# the score is clamped to.
#
# ATS_DOCUMENT_WEIGHT is shared equally across REQUIRED_DOCUMENTS, so the value
# of a single required document is ATS_DOCUMENT_WEIGHT / ATS_REQUIRED_DOC_COUNT.
# That was 60/9 when the required set had nine documents and is 60/10 now that
# TRANSCRIPT and GOOD_CONDUCT replaced KRA_PIN. Any change to the required set
# therefore invalidates every previously stored ats_score, which is persisted
# at submission time and never recomputed on read.
#
# "Academic Standing / Grade Award" was removed from the applicant biodata. It
# was never one of these criteria -- the keyword corpus reads institution,
# qualification, field of study, experience, trainings and document synonyms,
# and never included grade_award -- so removing it required no rebalancing and
# changed no applicant's score. The weights below are unchanged from before
# that field was dropped.
ATS_DOCUMENT_WEIGHT = 60.0
ATS_KEYWORD_WEIGHT = 40.0
ATS_SCORE_MAX = ATS_DOCUMENT_WEIGHT + ATS_KEYWORD_WEIGHT  # 100.0


class ApplyThrottle(UserRateThrottle):
    scope = "apply"


# ---------------------------------------------------------------------------
# Department scope resolution (single source of truth for Director isolation)
# ---------------------------------------------------------------------------
#
# The two "department level" roles behave identically everywhere in this API.
# They are treated as one Director role for every scoping decision below.

DIRECTOR_ROLES = ("DEPARTMENT_DIRECTOR", "DEPARTMENT")
CROSS_DEPARTMENT_ROLES = ("ADMIN", "HR")


def is_director_role(user):
    """True when the user is scoped to a single department."""
    return bool(
        user
        and user.is_authenticated
        and getattr(user, "role", None) in DIRECTOR_ROLES
    )


def get_director_department(user):
    """Resolve the department a Director is bound to, server-side only.

    Mirrors the resolution order already established by the Requisition
    endpoints (the reference pattern for Director isolation): a department the
    user explicitly directs takes precedence over the department recorded on
    their user account.  Returns ``None`` when the Director has no department
    assigned, in which case callers must return an empty result set rather than
    an unscoped one.
    """
    if not user or not user.is_authenticated:
        return None
    department = getattr(user, "directed_department", None)
    if department is None:
        department = getattr(user, "department", None)
    return department


def resolve_requested_department(request):
    """Resolve the department a scoped request must be limited to.

    Admin/HR are institutional roles: they may target any department through the
    ``department`` query parameter, or omit it entirely for a system-wide
    result.

    Directors can ONLY ever be scoped to their own department, and that
    department is always derived from the authenticated user — it is never read
    from client input.  A Director who supplies a ``department`` value that is
    not their own is rejected with 403 rather than having the value trusted or
    silently ignored, matching the convention already used when a Director
    submits a Requisition for another department.
    """
    user = request.user
    requested = request.query_params.get("department")

    if not is_director_role(user):
        if requested and requested != "ALL":
            return Department.objects.filter(id=requested).first()
        return None

    director_department = get_director_department(user)
    if requested and str(getattr(director_department, "id", "")) != str(requested):
        raise PermissionDenied(
            "You can only request data for your own assigned department."
        )
    return director_department


def scope_queryset_to_request_department(request, queryset, department_lookup):
    """Apply the resolved department scope to ``queryset``.

    ``department_lookup`` is the ORM path from the queryset's model to the
    owning Department, e.g. ``"job__department"`` or ``"department"``.

    A Director with no department assigned yields an empty queryset — never an
    unscoped one.
    """
    department = resolve_requested_department(request)

    if is_director_role(request.user) and department is None:
        return queryset.none()

    if department is not None:
        return queryset.filter(**{f"{department_lookup}_id": department.id})
    return queryset


# ---------------------------------------------------------------------------
# Reusable Permission Classes
# ---------------------------------------------------------------------------

class IsAdminOnly(permissions.BasePermission):
    def has_permission(self, request, view):
        return bool(request.user and request.user.is_authenticated and request.user.role == "ADMIN")


class IsHRorAdmin(permissions.BasePermission):
    def has_permission(self, request, view):
        return bool(request.user and request.user.is_authenticated and request.user.role in ["ADMIN", "HR"])


class IsDepartmentOrHRorAdmin(permissions.BasePermission):
    def has_permission(self, request, view):
        return bool(request.user and request.user.is_authenticated and request.user.role in ["ADMIN", "HR", "DEPARTMENT_DIRECTOR", "DEPARTMENT"])


class IsDirectorOrHRorAdmin(permissions.BasePermission):
    def has_permission(self, request, view):
        return bool(request.user and request.user.is_authenticated and request.user.role in ["ADMIN", "HR", "DEPARTMENT_DIRECTOR", "DEPARTMENT"])


class IsDepartmentDirectorOnly(permissions.BasePermission):
    def has_permission(self, request, view):
        return bool(request.user and request.user.is_authenticated and request.user.role in ["DEPARTMENT_DIRECTOR", "DEPARTMENT"])


# ---------------------------------------------------------------------------
# Recruitment authority: HR owns the whole workflow, Admin has read-only
# visibility, and only departments may raise requisitions.
#
# These replace the previous ``IsManagementOrReadOnly`` class, which was used
# for every vacancy endpoint and the applicant status endpoint. It was unsuitable
# for two independent reasons:
#
#   1. It returned True on SAFE_METHODS before any role check, so it granted
#      read access to anonymous callers. That is only safe on an endpoint whose
#      queryset is itself department-scoped.
#   2. It granted writes to both ADMIN and HR, which is how Admin gained the
#      ability to post vacancies and record applicant outcomes.
# ---------------------------------------------------------------------------


class IsHROnly(permissions.BasePermission):
    """
    Write authority reserved to the HR role.

    Used for every operation that commits a recruitment outcome: posting or
    editing a vacancy, archiving one, recording an applicant's decision, and
    reviewing a departmental requisition. Admin is deliberately excluded -- Admin
    configures the system, HR runs the recruitment.
    """

    def has_permission(self, request, view):
        return bool(
            request.user
            and request.user.is_authenticated
            and request.user.role == "HR"
        )

    def has_object_permission(self, request, view, obj):
        return self.has_permission(request, view)


class IsHROrReadOnlyManagement(permissions.BasePermission):
    """
    Read for the management roles, write for HR alone.

    Unlike the class it replaces, this requires authentication before granting
    read access, so it is safe for endpoints whose queryset is not itself
    department-scoped.
    """

    READABLE_ROLES = ("ADMIN", "HR", "DEPARTMENT_DIRECTOR", "DEPARTMENT")

    def has_permission(self, request, view):
        user = request.user
        if not (user and user.is_authenticated):
            return False
        if request.method in permissions.SAFE_METHODS:
            return user.role in self.READABLE_ROLES
        return user.role == "HR"


class IsVacancyListReadable(IsHROrReadOnlyManagement):
    """
    Read access to the vacancy list for management roles *and* applicants.

    Applicants need this: the portal's Vacancies page is where they discover
    openings and apply. They were refused with a 403, the frontend's catch
    handler cleared the list, and the page rendered "no vacancies" -- so the
    portal silently showed an applicant an empty system while the same
    vacancies were visible on the public landing page, which reads a different
    endpoint. Nothing surfaced the 403 anywhere.

    Write remains HR-only. Read scope is not widened here -- the view's
    ``get_queryset`` restricts applicants to open, applicable vacancies, and
    archived vacancies stay management-only records.
    """

    READABLE_ROLES = IsHROrReadOnlyManagement.READABLE_ROLES + ("APPLICANT",)


class IsVacancyManager(IsHROrReadOnlyManagement):
    """
    Vacancy management rights: HR edits, Admin may additionally delete.

    Deletion is separated from editing on purpose. HR owns posting and editing a
    vacancy, but an Admin also has to be able to withdraw one -- previously the
    delete button in the Vacancies page returned 403 for an Admin and surfaced as
    a bare "Failed to delete job" alert. Granting Admin full write instead would
    quietly hand over authoring rights that are not the Admin's.

    Read behaviour is inherited unchanged.
    """

    DELETE_ROLES = ("ADMIN", "HR")

    def has_permission(self, request, view):
        if not (request.user and request.user.is_authenticated):
            return False
        if request.method == "DELETE":
            return request.user.role in self.DELETE_ROLES
        return super().has_permission(request, view)


class IsDepartmentSubmitter(permissions.BasePermission):
    """
    Departments raise requisitions; HR receives and reviews them.

    Read is open to any authenticated user (the view's ``get_queryset`` applies
    the department scoping), but only DEPARTMENT_DIRECTOR / DEPARTMENT may
    *create*. HR and Admin are excluded from creation so a requisition always
    originates from a department requesting attachees and is reviewed by
    somebody other than its author.
    """

    def has_permission(self, request, view):
        user = request.user
        if not (user and user.is_authenticated):
            return False
        if request.method in permissions.SAFE_METHODS:
            return True
        return user.role in DIRECTOR_ROLES


# ---------------------------------------------------------------------------
# 1.1 Department Management Views (Full CRUD, Soft-Delete, Admin-Write)
# ---------------------------------------------------------------------------

class AvailableDirectorsListView(APIView):
    permission_classes = [IsAdminOnly]

    def get(self, request):
        directors = User.objects.filter(
            role__in=["DEPARTMENT_DIRECTOR", "DEPARTMENT"]
        ).select_related("department", "directed_department").order_by("first_name", "username")
        results = []
        for d in directors:
            assigned_dept = getattr(d, "directed_department", None) or d.department
            results.append({
                "id": d.id,
                "username": d.username,
                "email": d.email,
                "first_name": d.first_name,
                "last_name": d.last_name,
                "role": d.role,
                "assigned_department_id": assigned_dept.id if assigned_dept else None,
                "assigned_department_name": assigned_dept.name if assigned_dept else None,
            })
        return Response(results, status=status.HTTP_200_OK)


class DepartmentListCreateView(generics.ListCreateAPIView):
    serializer_class = DepartmentSerializer

    def get_permissions(self):
        if self.request.method == "POST":
            return [IsAdminOnly()]
        return [permissions.AllowAny()]

    def get_queryset(self):
        from django.db.models import Count, Sum, Q
        qs = Department.objects.select_related("director").annotate(
            annotated_active_vacancies=Count("jobs", filter=Q(jobs__is_active=True, jobs__is_archived=False), distinct=True),
            annotated_total_slots=Sum("jobs__slots_required", filter=Q(jobs__is_active=True, jobs__is_archived=False))
        ).order_by("name")
        # Anonymous users or regular applicants see active departments
        if not self.request.user.is_authenticated or self.request.user.role == "APPLICANT":
            qs = qs.filter(is_active=True)
        return qs

    def create(self, request, *args, **kwargs):
        from accounts.audit import log_audit_event
        from django.contrib.auth import password_validation
        from django.core.exceptions import ValidationError as DjangoValidationError
        import secrets
        import string

        data = request.data.copy()
        new_director_data = data.pop("new_director", None)
        director_user = None

        if new_director_data and isinstance(new_director_data, dict):
            email = new_director_data.get("email", "").strip().lower()
            username = new_director_data.get("username", "").strip() or email.split("@")[0]
            password = new_director_data.get("password", "").strip()
            first_name = new_director_data.get("first_name", "").strip()
            last_name = new_director_data.get("last_name", "").strip()

            if not email:
                return Response({"detail": "Director email is required."}, status=status.HTTP_400_BAD_REQUEST)
            if User.objects.filter(email__iexact=email).exists() or User.objects.filter(username=username).exists():
                return Response({"detail": "A user with this email or username already exists."}, status=status.HTTP_400_BAD_REQUEST)

            if not password:
                alphabet = string.ascii_letters + string.digits + "!@#$%^&*"
                password = "".join(secrets.choice(alphabet) for _ in range(14))

            try:
                password_validation.validate_password(password)
            except DjangoValidationError as err:
                return Response({"detail": f"Director password does not meet security requirements: {', '.join(err.messages)}"}, status=status.HTTP_400_BAD_REQUEST)

            director_user = User.objects.create_user(
                username=username,
                email=email,
                password=password,
                first_name=first_name,
                last_name=last_name,
                role="DEPARTMENT_DIRECTOR",
                must_change_password=True,
            )
            data["director"] = director_user.id

        serializer = self.get_serializer(data=data)
        serializer.is_valid(raise_exception=True)
        dept = serializer.save()

        if director_user:
            director_user.department = dept
            director_user.save()
        elif dept.director:
            dept.director.department = dept
            dept.director.save()

        log_audit_event(
            request,
            "DEPARTMENT_CREATED",
            f"Created department '{dept.name}'",
            metadata={"department_id": dept.id, "department_name": dept.name, "director_id": dept.director_id}
        )

        headers = self.get_success_headers(serializer.data)
        return Response(DepartmentSerializer(dept).data, status=status.HTTP_201_CREATED, headers=headers)


class DepartmentDetailView(generics.RetrieveUpdateDestroyAPIView):
    queryset = Department.objects.select_related("director").all()
    serializer_class = DepartmentSerializer

    def get_permissions(self):
        if self.request.method in permissions.SAFE_METHODS:
            return [permissions.AllowAny()]
        return [IsAdminOnly()]

    def update(self, request, *args, **kwargs):
        from accounts.audit import log_audit_event
        from django.contrib.auth import password_validation
        from django.core.exceptions import ValidationError as DjangoValidationError
        import secrets
        import string

        partial = kwargs.pop("partial", False)
        instance = self.get_object()
        data = request.data.copy()
        new_director_data = data.pop("new_director", None)
        director_user = None

        if new_director_data and isinstance(new_director_data, dict):
            email = new_director_data.get("email", "").strip().lower()
            username = new_director_data.get("username", "").strip() or email.split("@")[0]
            password = new_director_data.get("password", "").strip()
            first_name = new_director_data.get("first_name", "").strip()
            last_name = new_director_data.get("last_name", "").strip()

            if not email:
                return Response({"detail": "Director email is required."}, status=status.HTTP_400_BAD_REQUEST)
            if User.objects.filter(email__iexact=email).exists() or User.objects.filter(username=username).exists():
                return Response({"detail": "A user with this email or username already exists."}, status=status.HTTP_400_BAD_REQUEST)

            if not password:
                alphabet = string.ascii_letters + string.digits + "!@#$%^&*"
                password = "".join(secrets.choice(alphabet) for _ in range(14))

            try:
                password_validation.validate_password(password)
            except DjangoValidationError as err:
                return Response({"detail": f"Director password does not meet security requirements: {', '.join(err.messages)}"}, status=status.HTTP_400_BAD_REQUEST)

            director_user = User.objects.create_user(
                username=username,
                email=email,
                password=password,
                first_name=first_name,
                last_name=last_name,
                role="DEPARTMENT_DIRECTOR",
                must_change_password=True,
            )
            data["director"] = director_user.id

        serializer = self.get_serializer(instance, data=data, partial=partial)
        serializer.is_valid(raise_exception=True)
        dept = serializer.save()

        if director_user:
            director_user.department = dept
            director_user.save()
        elif dept.director:
            dept.director.department = dept
            dept.director.save()

        log_audit_event(
            request,
            "DEPARTMENT_UPDATED",
            f"Updated department '{dept.name}'",
            metadata={"department_id": dept.id, "department_name": dept.name, "director_id": dept.director_id}
        )

        return Response(DepartmentSerializer(dept).data)

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        
        # If explicitly requested to deactivate rather than hard-delete
        if request.query_params.get("deactivate") == "true":
            instance.is_active = False
            instance.save(update_fields=["is_active", "updated_at"])
            return Response(
                {"detail": f"Department '{instance.name}' has been deactivated successfully.", "is_active": False},
                status=status.HTTP_200_OK
            )

        # Check for dependent records
        jobs_count = instance.jobs.count()
        reqs_count = instance.requisitions.count()
        deps_count = instance.deployments.count() if hasattr(instance, "deployments") else 0
        staff_count = instance.staff_users.count() if hasattr(instance, "staff_users") else 0
        notifs_count = instance.notifications.count() if hasattr(instance, "notifications") else 0
        total_dependencies = jobs_count + reqs_count + deps_count + staff_count + notifs_count

        if total_dependencies > 0:
            dep_details = []
            if jobs_count: dep_details.append(f"{jobs_count} attachment vacancies")
            if reqs_count: dep_details.append(f"{reqs_count} requisitions")
            if deps_count: dep_details.append(f"{deps_count} student deployments")
            if staff_count: dep_details.append(f"{staff_count} assigned staff users")
            if notifs_count: dep_details.append(f"{notifs_count} notifications")
            dep_str = ", ".join(dep_details)

            return Response(
                {
                    "detail": f"Cannot permanently delete '{instance.name}' because it is referenced by {dep_str}. Deactivating this department is recommended to safely preserve institutional records.",
                    "has_dependencies": True,
                    "can_deactivate": True,
                    "dependencies": {
                        "jobs": jobs_count,
                        "requisitions": reqs_count,
                        "deployments": deps_count,
                        "staff_users": staff_count,
                        "notifications": notifs_count,
                    }
                },
                status=status.HTTP_400_BAD_REQUEST
            )

        # If a director is assigned to this department, unlink it before deletion
        if instance.director:
            instance.director = None
            instance.save(update_fields=["director"])

        # No dependencies exist: genuine hard delete
        instance.delete()
        return Response(
            {"detail": f"Department '{instance.name}' was permanently deleted."},
            status=status.HTTP_204_NO_CONTENT
        )


# ---------------------------------------------------------------------------
# 2. Requisition Management Views (Department -> HR -> Auto-Vacancy)
# ---------------------------------------------------------------------------

class RequisitionListCreateView(generics.ListCreateAPIView):
    """
    Departments raise requisitions for attachees; only HR receives and reviews them.

    Creation is restricted to DEPARTMENT_DIRECTOR / DEPARTMENT via
    ``IsDepartmentSubmitter``, so a requisition always originates from a
    department that actually needs the placements. HR and Admin can read every
    requisition but cannot author one, which also removes the possibility of an
    HR user approving a requisition they wrote themselves.
    """

    from .serializers import RequisitionSerializer
    serializer_class = RequisitionSerializer
    permission_classes = (IsDepartmentSubmitter,)

    def get_queryset(self):
        user = self.request.user
        qs = Requisition.objects.select_related(
            "department", "requested_by", "reviewed_by", "created_job"
        ).order_by("-created_at")

        if user.role in ["ADMIN", "HR"]:
            dept_id = self.request.query_params.get("department")
            if dept_id and dept_id != "ALL":
                if not dept_id.isdigit():
                    return Requisition.objects.none()
                qs = qs.filter(department_id=dept_id)
            status_param = self.request.query_params.get("status")
            if status_param:
                qs = qs.filter(status=status_param)
            return qs

        elif is_director_role(user):
            # Strict scoping: Director can only view their own department's requisitions
            director_dept = get_director_department(user)
            if director_dept:
                return qs.filter(Q(department=director_dept) | Q(requested_by=user))
            return qs.filter(requested_by=user)

        # Applicants cannot view internal departmental requisitions
        return Requisition.objects.none()

    def create(self, request, *args, **kwargs):
        user = request.user
        data = request.data.copy()

        # The department is always resolved server-side from the submitting
        # Director. ``IsDepartmentSubmitter`` has already rejected every other
        # role, so by the time this runs the caller is guaranteed to be a
        # department user, and the two error branches are reachable states
        # rather than dead code.
        director_dept = get_director_department(user)
        if not director_dept:
            return Response(
                {"detail": "You do not have an assigned department to submit requisitions for."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # Guarded comparison: a non-numeric "department" previously reached
        # int() unguarded and raised ValueError, which surfaced as a 500 instead
        # of a client error.
        dept_id = data.get("department")
        if dept_id not in (None, ""):
            if not str(dept_id).isdigit():
                return Response(
                    {"detail": "'department' must be a numeric department id."},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            if int(dept_id) != director_dept.id:
                return Response(
                    {"detail": "You can only submit requisitions for your own assigned department."},
                    status=status.HTTP_403_FORBIDDEN,
                )
        data["department"] = director_dept.id

        serializer = self.get_serializer(data=data)
        serializer.is_valid(raise_exception=True)
        req_obj = serializer.save(requested_by=user, status="PENDING")

        # Route the notification to the HR queue rather than broadcasting it to
        # the department inbox. A requisition is addressed to HR, and a
        # department-scoped notification with no recipient is never read by the
        # people who have to act on it.
        hr_users = User.objects.filter(role="HR", is_active=True)
        for hr_user in hr_users:
            Notification.objects.create(
                recipient=hr_user,
                department=req_obj.department,
                title="New Attachment Requisition Submitted",
                message=(
                    f"{user.first_name} {user.last_name} ({req_obj.department.name}) "
                    f"submitted a requisition for "
                    f"{req_obj.num_candidates_requested} candidate(s), awaiting HR review."
                ),
                notification_type="GENERAL",
            )

        return Response(RequisitionSerializer(req_obj).data, status=status.HTTP_201_CREATED)


class RequisitionDetailView(generics.RetrieveAPIView):
    from .serializers import RequisitionSerializer
    queryset = Requisition.objects.select_related(
        "department", "requested_by", "reviewed_by", "created_job"
    ).all()
    serializer_class = RequisitionSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_object(self):
        obj = super().get_object()
        user = self.request.user
        if user.role in ["ADMIN", "HR"]:
            return obj
        if is_director_role(user):
            director_dept = get_director_department(user)
            if (director_dept and obj.department_id == director_dept.id) or obj.requested_by_id == user.id:
                return obj
            raise permissions.exceptions.PermissionDenied("You do not have access to view this department's requisition.")
        raise permissions.exceptions.PermissionDenied("Access denied.")


class RequisitionReviewView(APIView):
    """
    HR reviews a departmental requisition and, on approval, publishes the vacancy.

    Restricted to the HR role. Admin is deliberately excluded: requisitions route
    from a department to HR, and adding Admin here would put a second role in
    the path between a department's request and the vacancies that get posted.
    """

    permission_classes = [IsHROnly]

    def post(self, request, pk):
        try:
            req_obj = Requisition.objects.select_related("department", "requested_by").get(pk=pk)
        except Requisition.DoesNotExist:
            return Response({"detail": "Requisition not found."}, status=status.HTTP_404_NOT_FOUND)

        action = request.data.get("action", "").upper()
        review_notes = request.data.get("review_notes", "").strip()

        if action not in ("APPROVE", "REJECT"):
            return Response(
                {"detail": "Invalid action. Must be 'APPROVE' or 'REJECT'."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # A decision is recorded once. Previously there was no status guard, so
        # re-POSTing APPROVE on an already-approved requisition created a second
        # live vacancy and silently repointed created_job at it, leaving two
        # published vacancies for one request and an orphaned first vacancy.
        if req_obj.status != "PENDING":
            return Response(
                {
                    "detail": (
                        f"This requisition has already been "
                        f"{req_obj.status.lower()} and cannot be reviewed again."
                    ),
                    "status": req_obj.status,
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        # Belt-and-braces guard. IsDepartmentSubmitter already prevents HR from
        # authoring a requisition, so this cannot currently fire, but it keeps
        # the rule explicit and survives a future change to who may submit.
        if req_obj.requested_by_id == request.user.id:
            return Response(
                {"detail": "You cannot review a requisition that you submitted."},
                status=status.HTTP_403_FORBIDDEN,
            )

        if action == "APPROVE":
            req_obj.status = "APPROVED"
            req_obj.reviewed_by = request.user
            req_obj.reviewed_at = timezone.now()
            req_obj.review_notes = review_notes

            # Derive deadline from duration_start or default 30 days
            deadline_dt = timezone.datetime.combine(req_obj.duration_start, timezone.datetime.min.time())
            if timezone.is_naive(deadline_dt):
                deadline_dt = timezone.make_aware(deadline_dt)
            if deadline_dt <= timezone.now():
                deadline_dt = timezone.now() + timedelta(days=30)

            job_title = request.data.get("job_title") or f"Industrial Attachee - {req_obj.department.name}"
            
            # Auto-create Job vacancy linked to this requisition
            job = Job.objects.create(
                department=req_obj.department,
                title=job_title,
                description=f"Industrial Attachment opportunity in the {req_obj.department.name}. Successful candidates will receive structured on-the-job training, technical mentorship, and practical public sector experience.",
                requirements=req_obj.requirements,
                job_type="ATTACHMENT",
                location=req_obj.department.location or "State Department for Petroleum HQ",
                slots_required=req_obj.num_candidates_requested,
                deadline=deadline_dt,
                is_active=True,
                is_archived=False,
            )

            req_obj.created_job = job
            req_obj.save()

            # Trigger notification
            Notification.objects.create(
                department=req_obj.department,
                recipient=req_obj.requested_by,
                title=f"Requisition Approved: {job.title}",
                message=f"Requisition #{req_obj.id} for {req_obj.num_candidates_requested} candidate(s) has been approved by HR and published as a vacancy.",
                notification_type="JOB_POSTED",
            )

            from .serializers import RequisitionSerializer
            return Response(
                {
                    "detail": "Requisition approved and attachment vacancy generated successfully.",
                    "requisition": RequisitionSerializer(req_obj).data,
                    "job_id": job.id,
                },
                status=status.HTTP_200_OK,
            )

        if action == "REJECT":
            # The condition was previously written as "APPROVE", duplicating the
            # branch above. It behaved correctly only because that branch
            # returns first, so this was reached for REJECT and nothing else --
            # but the label claimed the opposite and invited a future edit that
            # would have deleted the rejection path.
            req_obj.status = "REJECTED"
            req_obj.reviewed_by = request.user
            req_obj.reviewed_at = timezone.now()
            req_obj.review_notes = review_notes
            req_obj.save()

            Notification.objects.create(
                department=req_obj.department,
                recipient=req_obj.requested_by,
                title=f"Requisition Rejected: {req_obj.department.name}",
                message=f"Your requisition #{req_obj.id} was not approved. Feedback: {review_notes or 'No feedback provided'}",
                notification_type="GENERAL",
            )

            from .serializers import RequisitionSerializer
            return Response(
                {
                    "detail": "Requisition marked as rejected.",
                    "requisition": RequisitionSerializer(req_obj).data,
                },
                status=status.HTTP_200_OK,
            )


# ---------------------------------------------------------------------------
# 1.3 Public Vacancies View (Unauthenticated, AllowAny)
# ---------------------------------------------------------------------------

class PublicVacancyListView(generics.ListAPIView):
    serializer_class = PublicJobSerializer
    permission_classes = [permissions.AllowAny]

    def get_queryset(self):
        now = timezone.now()
        qs = Job.objects.filter(
            is_active=True,
            is_archived=False,
            deadline__gte=now,
            job_type="ATTACHMENT"
        ).select_related("department").order_by("-created_at")

        department_id = self.request.query_params.get("department")
        if department_id:
            qs = qs.filter(department_id=department_id)

        limit = self.request.query_params.get("limit")
        if limit and limit.isdigit():
            qs = qs[:int(limit)]
        return qs

    def list(self, request, *args, **kwargs):
        from django.core.cache import cache
        from .signals import CACHE_KEY_PUBLIC_VACANCIES, CACHE_TTL_PUBLIC_VACANCIES

        # Check cache if requesting standard public vacancies list without query filters
        if not request.query_params:
            cached_data = cache.get(CACHE_KEY_PUBLIC_VACANCIES)
            if cached_data is not None:
                return Response(cached_data)

            response = super().list(request, *args, **kwargs)
            cache.set(
                CACHE_KEY_PUBLIC_VACANCIES,
                response.data,
                timeout=self._public_vacancy_cache_ttl(),
            )
            return response

        return super().list(request, *args, **kwargs)

    @staticmethod
    def _public_vacancy_cache_ttl(ceiling=None):
        """
        Never cache a vacancy list for longer than its soonest deadline.

        The list is filtered on ``deadline__gte=now``, so its membership is
        time-dependent: a vacancy closes on a schedule that no save touches.
        With a flat 15-minute TTL, a vacancy whose deadline fell one minute ago
        kept appearing on the public landing page for another fourteen, and the
        cache was only ever invalidated by a Job save, which closing a deadline
        is not.

        The TTL is therefore the shorter of the ceiling and the time remaining
        until the next deadline. The soonest deadline bounds how long the
        cached payload can still be correct, so an expiring entry stops being
        served at the moment it goes stale. An entry whose deadline is already
        past is not cached at all.
        """
        if ceiling is None:
            from .signals import CACHE_TTL_PUBLIC_VACANCIES

            ceiling = CACHE_TTL_PUBLIC_VACANCIES

        soonest = (
            Job.objects.filter(
                is_active=True, is_archived=False, job_type="ATTACHMENT"
            )
            .exclude(deadline__lt=timezone.now())
            .order_by("deadline")
            .values_list("deadline", flat=True)
            .first()
        )
        if soonest is None:
            # Nothing left to advertise. Caching the empty list is fine and
            # avoids rebuilding it on every public request.
            return ceiling

        seconds_until_expiry = (soonest - timezone.now()).total_seconds()
        if seconds_until_expiry <= 0:
            return 0
        return max(1, int(min(ceiling, seconds_until_expiry)))


# ---------------------------------------------------------------------------
# Vacancy & Archiving Views
# ---------------------------------------------------------------------------

class JobListCreateView(generics.ListCreateAPIView):
    serializer_class = JobSerializer
    # HR posts vacancies. Admin retains read-only visibility; nobody else may
    # write. Applicants may also read -- the portal's Vacancies page is where
    # they find openings -- and their queryset is narrowed below.
    permission_classes = (IsVacancyListReadable,)

    def get_queryset(self):
        qs = Job.objects.select_related("department").order_by("-created_at")

        is_applicant = self.request.user.role == "APPLICANT"
        if is_applicant:
            # An applicant is shown exactly what they could act on: an open,
            # unexpired, attachment-type posting. Without this the widened read
            # permission would also expose deactivated, closed and internal
            # postings, plus archived ones via the query params below.
            qs = qs.filter(
                is_active=True,
                is_archived=False,
                job_type="ATTACHMENT",
                deadline__gte=timezone.now(),
            )
            # Archived vacancies are internal records, so the archive toggles are
            # ignored for applicants rather than being honoured and then filtered.
            department_id = self.request.query_params.get("department")
            if department_id and department_id != "ALL":
                qs = qs.filter(department_id=department_id)
            return qs

        # Archived vacancies are internal records. They are browsable by
        # Admin/HR, and by Directors only through their own department-scoped
        # Archives module — never as a cross-department listing from here.
        include_archived = self.request.query_params.get("include_archived", "").lower() == "true"
        archived_only = self.request.query_params.get("archived", "").lower() == "true"

        if is_director_role(self.request.user) and (include_archived or archived_only):
            raise PermissionDenied(
                "Archived vacancies are only available through the department archives module."
            )

        if archived_only:
            qs = qs.filter(is_archived=True)
        elif not include_archived:
            qs = qs.filter(is_archived=False)

        department_id = self.request.query_params.get("department")
        if department_id and department_id != "ALL":
            qs = qs.filter(department_id=department_id)

        job_type = self.request.query_params.get("job_type")
        if job_type:
            qs = qs.filter(job_type=job_type)

        return qs

    def perform_create(self, serializer):
        job = serializer.save()

        # Trigger department notification when a vacancy is posted for a department
        if job.department:
            Notification.objects.create(
                department=job.department,
                title=f"New Vacancy Posted: {job.title}",
                message=(
                    f"A new vacancy '{job.title}' with {job.slots_required} slot(s) "
                    f"has been posted for {job.department.name}."
                ),
                notification_type="JOB_POSTED",
            )
            # If contact email exists, send safe email notification
            if job.department.contact_email:
                try:
                    from django.core.mail import send_mail
                    send_mail(
                        subject=f"New Vacancy Notice: {job.title}",
                        message=(
                            f"Greetings,\n\nA new vacancy '{job.title}' ({job.slots_required} slots) "
                            f"has been posted for your department ({job.department.name}).\n"
                            f"Application Deadline: {job.deadline}\n\n"
                            "State Department for Petroleum Industrial Attachment Portal"
                        ),
                        from_email=None,
                        recipient_list=[job.department.contact_email],
                        fail_silently=True,
                    )
                except Exception:
                    pass


class IsHROrAdmin(permissions.BasePermission):
    """
    HR and Admin only.

    Used for the record-management operations that withdraw a vacancy from the
    active list and put it back. Admin is included because Admin can delete a
    vacancy, and a delete that could not be undone by the same role would be a
    trap: the vacancy would be stuck in the archives with only HR able to
    recover it.
    """

    def has_permission(self, request, view):
        user = request.user
        return bool(user and user.is_authenticated) and user.role in ("HR", "ADMIN")


class JobArchiveToggleView(APIView):
    permission_classes = (IsHROrAdmin,)

    def post(self, request, pk):
        try:
            job = Job.objects.get(pk=pk)
        except Job.DoesNotExist:
            return Response({"detail": "Job not found."}, status=status.HTTP_404_NOT_FOUND)

        job.is_archived = not job.is_archived
        # Keep "archived" and "active" in agreement. Deletion and the System
        # Archives restore both maintain this, so leaving the toggle to flip only
        # `is_archived` produced vacancies whose listed state contradicted what
        # the public and applicant endpoints would serve.
        job.is_active = not job.is_archived
        if job.is_archived:
            job.archived_at = timezone.now()
            job.archived_by = request.user if request.user.is_authenticated else None
        else:
            job.archived_at = None
            job.archived_by = None
        job.save(update_fields=["is_archived", "is_active", "archived_at", "archived_by"])

        action = "archived" if job.is_archived else "unarchived"
        return Response(
            {
                "detail": f"Job '{job.title}' has been {action} successfully.",
                "is_archived": job.is_archived,
            },
            status=status.HTTP_200_OK,
        )


class JobBulkArchiveView(APIView):
    permission_classes = (IsHROnly,)

    def post(self, request):
        now = timezone.now()
        job_ids = request.data.get("job_ids", [])
        if job_ids and isinstance(job_ids, list):
            count = Job.objects.filter(id__in=job_ids, is_archived=False).update(is_archived=True)
            return Response(
                {
                    "detail": f"Successfully archived {count} selected vacancies.",
                    "archived_count": count,
                },
                status=status.HTTP_200_OK,
            )
        expired_jobs = Job.objects.filter(deadline__lt=now, is_archived=False)
        count = expired_jobs.update(is_archived=True)
        return Response(
            {
                "detail": f"Successfully archived {count} past-deadline vacancies.",
                "archived_count": count,
            },
            status=status.HTTP_200_OK,
        )


class JobDetailView(generics.RetrieveUpdateDestroyAPIView):
    queryset = Job.objects.all()
    serializer_class = JobSerializer
    # Editing is HR-only; read is open to the management roles; deletion is
    # additionally open to Admin. See IsVacancyManager.
    permission_classes = (IsVacancyManager,)

    def perform_destroy(self, instance):
        # Archive rather than hard-delete.
        #
        # `Application.job` is on_delete=CASCADE, so removing the row would take
        # the applicant's application with it, and from there the deployment,
        # clearance and recommendation letter -- irreversibly, and without
        # leaving anything for the audit trail or the System Archives module to
        # show. Archiving withdraws the vacancy from the active list (it is
        # flagged inactive, so it also disappears from the public endpoint) and
        # leaves it restorable.
        instance.is_archived = True
        instance.is_active = False
        instance.archived_at = timezone.now()
        instance.archived_by = self.request.user if self.request.user and self.request.user.is_authenticated else None
        instance.save(update_fields=["is_archived", "is_active", "archived_at", "archived_by"])
        from accounts.audit import log_audit_event
        log_audit_event(
            self.request,
            "RECORD_ARCHIVED",
            f"Archived vacancy '{instance.title}' (ID: {instance.id})",
            metadata={"entity_type": "VACANCY", "id": instance.id, "title": instance.title}
        )


# ---------------------------------------------------------------------------
# Applications & ATS Scoring
# ---------------------------------------------------------------------------

class ApplicationListCreateView(generics.ListCreateAPIView):
    serializer_class = ApplicationSerializer
    permission_classes = (permissions.IsAuthenticated,)
    throttle_classes = [ApplyThrottle]

    def get_queryset(self):
        user = self.request.user
        if user.role in ["ADMIN", "HR"]:
            qs = Application.objects.all().select_related("user", "job", "job__department").prefetch_related("user__documents", "status_events")
        elif is_director_role(user):
            # Same scoping as ApplicationDetailView: both department-level roles
            # only see applications raised against their own department's jobs.
            director_dept = get_director_department(user)
            if director_dept is None:
                return Application.objects.none()
            qs = Application.objects.filter(job__department=director_dept).select_related("user", "job", "job__department").prefetch_related("user__documents", "status_events")
        else:
            qs = Application.objects.filter(user=user).select_related("user", "job", "job__department").prefetch_related("user__documents", "status_events")

        job_type = self.request.query_params.get("job_type")
        if job_type:
            qs = qs.filter(job__job_type=job_type)
        return qs

    def perform_create(self, serializer):
        user = self.request.user
        job = serializer.validated_data["job"]

        if job.is_full:
            from rest_framework.exceptions import ValidationError
            raise ValidationError(
                {"detail": f"Applications for '{job.title}' are closed. All {job.slots_required} slot(s) have been filled."}
            )

        if timezone.now() > job.deadline:
            from rest_framework.exceptions import ValidationError
            raise ValidationError(
                {"detail": f"Applications for '{job.title}' are closed. The deadline has passed."}
            )

        # 1. Authoritative server-side check on Biodata completeness
        profile = getattr(user, "profile", None)
        missing_biodata = []
        if not user.first_name or not user.first_name.strip(): missing_biodata.append("First Name")
        if not user.last_name or not user.last_name.strip(): missing_biodata.append("Last Name")
        if not user.email or not user.email.strip(): missing_biodata.append("Email Address")
        if not profile:
            missing_biodata.extend(["Student Biodata Profile", "Phone Number", "ID Number", "Date of Birth", "Institution Name", "Qualification", "Field of Study", "Joining Date", "Next of Kin Details"])
        else:
            if not profile.phone_number or not profile.phone_number.strip(): missing_biodata.append("Phone Number")
            if not profile.id_number or not profile.id_number.strip(): missing_biodata.append("National ID Number")
            if not profile.dob: missing_biodata.append("Date of Birth")
            if not profile.gender: missing_biodata.append("Gender")
            if not profile.nationality or not profile.nationality.strip(): missing_biodata.append("Nationality")
            if not profile.postal_address or not profile.postal_address.strip(): missing_biodata.append("Postal Address")
            if not profile.institution_name or not profile.institution_name.strip(): missing_biodata.append("Institution / University Name")
            if not profile.qualification or not profile.qualification.strip(): missing_biodata.append("Pursued Qualification")
            if not profile.field_of_study or not profile.field_of_study.strip(): missing_biodata.append("Field of Study")
            if not profile.joining_date: missing_biodata.append("Intended Attachment Start Date")
            if not profile.next_of_kin_name or not profile.next_of_kin_name.strip(): missing_biodata.append("Next of Kin Name")
            if not profile.next_of_kin_phone or not profile.next_of_kin_phone.strip(): missing_biodata.append("Next of Kin Phone Number")

        if missing_biodata:
            from rest_framework.exceptions import ValidationError
            raise ValidationError(
                {"detail": f"Cannot submit application. Required student biodata is incomplete: {', '.join(missing_biodata)}."}
            )

        # 2. Authoritative server-side check on the 10 Required Verification Documents
        user_docs = set(user.documents.values_list("document_type", flat=True))
        missing_docs = [doc for doc in REQUIRED_DOCUMENTS if doc not in user_docs]
        if missing_docs:
            from rest_framework.exceptions import ValidationError
            formatted_missing = [d.replace("_", " ").title() for d in missing_docs]
            raise ValidationError(
                {"detail": f"Cannot submit application. Missing required verification document(s): {', '.join(formatted_missing)}."}
            )

        ats_score = self.calculate_ats_score(self.request.user, job)
        application = serializer.save(user=self.request.user, ats_score=ats_score)

        # Open the applicant's journey with a "submitted" step. Without it the
        # timeline would begin at the first HR action, leaving the applicant with
        # no evidence that their own submission was ever received.
        ApplicationStatusEvent.objects.create(
            application=application,
            from_status="",
            to_status=application.status,
            changed_by=self.request.user,
            note="Application submitted",
        )

    def calculate_ats_score(self, user, job):
        """
        Scores an applicant 0-100 from two weighted criteria:

            documents = (uploaded required docs / total) * ATS_DOCUMENT_WEIGHT
            keywords  = (matched job words / total job words)
                        * ATS_KEYWORD_WEIGHT
            total     = clamp(documents + keywords, 0, ATS_SCORE_MAX)

        The weights sum to exactly ATS_SCORE_MAX, so a fully-documented applicant
        who matches every job keyword still tops out at 100 rather than leaving
        a gap. Academic Standing / Grade Award is deliberately not a criterion:
        the field has been removed from the biodata and never scored anything.
        """
        user_docs = set(user.documents.values_list("document_type", flat=True))
        uploaded_req_count = sum(1 for doc in REQUIRED_DOCUMENTS if doc in user_docs)
        doc_score = (uploaded_req_count / ATS_REQUIRED_DOC_COUNT) * ATS_DOCUMENT_WEIGHT

        # The vacancy description is Admin-authored rich text, so it is reduced
        # to plain words before matching. Without this, every tag name and
        # attribute in the markup ("strong", "href", "class", "http") would
        # enter the job keyword set, can never match applicant documents, and
        # would depress the keyword ratio -- and therefore every score.
        job_text = " ".join(
            rich_text_to_plain_text(part)
            for part in (job.title, job.description, job.requirements)
        ).lower()
        job_words = set(re.findall(r"\b[a-z]{3,}\b", job_text)) - ENGLISH_STOP_WORDS

        user_content_parts = []
        profile = getattr(user, "profile", None)
        if profile:
            if profile.institution_name:
                user_content_parts.append(profile.institution_name)
            if profile.qualification:
                user_content_parts.append(profile.qualification)
            if profile.field_of_study:
                user_content_parts.append(profile.field_of_study)
        if hasattr(user, "education"):
            for edu in user.education.all():
                user_content_parts.extend([edu.qualification, edu.field_of_study, edu.institution_name])
        if hasattr(user, "experience"):
            for exp in user.experience.all():
                user_content_parts.extend([exp.job_title, exp.responsibilities])
        if hasattr(user, "trainings"):
            for t in user.trainings.all():
                user_content_parts.append(t.name)
        for doc_type in user_docs:
            if doc_type in DOCUMENT_ATS_KEYWORDS:
                user_content_parts.append(DOCUMENT_ATS_KEYWORDS[doc_type])

        user_text = " ".join(user_content_parts).lower()
        user_words = set(re.findall(r"\b[a-z]{3,}\b", user_text)) - ENGLISH_STOP_WORDS

        if job_words:
            matched_words = job_words.intersection(user_words)
            keyword_score = (len(matched_words) / len(job_words)) * ATS_KEYWORD_WEIGHT
        else:
            keyword_score = ATS_KEYWORD_WEIGHT

        total_score = doc_score + keyword_score
        return round(min(max(total_score, 0.0), ATS_SCORE_MAX), 2)


class ApplicationDetailView(generics.RetrieveAPIView):
    queryset = Application.objects.all()
    serializer_class = ApplicationSerializer
    permission_classes = (permissions.IsAuthenticated,)

    def get_queryset(self):
        user = self.request.user
        if user.role in ["ADMIN", "HR"]:
            return Application.objects.all().select_related("user", "job", "job__department")
        elif is_director_role(user):
            dept = get_director_department(user)
            return Application.objects.filter(job__department=dept).select_related("user", "job", "job__department") if dept else Application.objects.none()
        return Application.objects.filter(user=user).select_related("user", "job", "job__department")


class ApplicationStatusUpdateView(APIView):
    """
    Records HR's decision on an applicant.

    Selection is binary. A decision is exactly one of:

      * ``SUCCESSFUL`` -- the opportunity is granted to the applicant;
      * ``REJECTED``  -- the application is declined.

    ``PENDING`` and ``REVIEWED`` are workflow markers, not outcomes, so they are
    accepted only as a reversal of a previous decision and never as a decision
    in their own right. Previously this endpoint accepted any of the five
    statuses including the intermediate SHORTLISTED and HIRED, and was open to
    Admin as well as HR.

    Only HR may call it (``IsHROnly``). Admin deliberately has no authority over
    applicant outcomes.
    """

    permission_classes = (IsHROnly,)

    def patch(self, request, pk):
        try:
            application = Application.objects.select_related("job", "job__department", "user").get(pk=pk)
        except Application.DoesNotExist:
            return Response({"detail": "Application not found."}, status=status.HTTP_404_NOT_FOUND)

        new_status = request.data.get("status")
        valid_statuses = [c[0] for c in Application.STATUS_CHOICES]
        if new_status not in valid_statuses:
            return Response(
                {
                    "detail": (
                        f"Invalid status value. Must be one of: "
                        f"{', '.join(valid_statuses)}."
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        is_decision = new_status in Application.DECISION_STATUSES
        was_decided = application.status in Application.DECISION_STATUSES

        # A workflow marker may only be used to reopen a decided application, so
        # an application can never be parked in PENDING/REVIEWED as if that were
        # an answer. This is what stops "no decision" from being mistaken for a
        # third outcome.
        if not is_decision and not was_decided:
            return Response(
                {
                    "detail": (
                        f"'{new_status}' is a workflow marker, not a selection "
                        f"outcome. An application must be decided as either "
                        f"{' or '.join(Application.DECISION_STATUSES)}."
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        # Clearing a decision is allowed so HR can correct a mistake, but it is
        # an explicit reversal rather than a decision.
        if not is_decision and application.status == "REVIEWED" and new_status == "PENDING":
            return Response(
                {"detail": "Application is already marked as REVIEWED and cannot be moved back to PENDING."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        previous_status = application.status
        application.status = new_status
        application.save()

        # Append the step to the persisted journey. Written only when the status
        # actually moved, so a repeated save of the same status does not pad the
        # applicant's timeline with duplicate entries. A failure to record the
        # event must not lose the decision itself, so it is isolated: the status
        # is already committed and the applicant is still notified below.
        if previous_status != new_status:
            try:
                ApplicationStatusEvent.objects.create(
                    application=application,
                    from_status=previous_status,
                    to_status=new_status,
                    changed_by=request.user,
                )
            except Exception:  # pragma: no cover - defensive
                logger.exception(
                    "Could not record status event for application %s", application.pk
                )

        # Notify only on a genuine decision, and only when the decision actually
        # changed, so re-saving the same status does not spam the applicant.
        if is_decision and previous_status != new_status:
            if new_status == "SUCCESSFUL":
                department_name = (
                    application.job.department.name
                    if application.job.department
                    else "the State Department for Petroleum"
                )
                Notification.objects.create(
                    recipient=application.user,
                    department=application.job.department,
                    title="Congratulations! Application Successful",
                    message=(
                        f"Your application for the industrial attachment "
                        f"opportunity '{application.job.title}' at "
                        f"{department_name} was successful. You have been "
                        f"granted the opportunity."
                    ),
                    notification_type="APPLICATION_UPDATE",
                    # Addressed to this applicant alone. Without it the notice
                    # also reached the HR/Admin feed (which is unfiltered) and the
                    # Director of the vacancy's department, both of whom have no
                    # need to read one applicant's result.
                    audience="APPLICANT",
                )
            else:
                Notification.objects.create(
                    recipient=application.user,
                    department=application.job.department,
                    title="Application Outcome: Not Successful",
                    message=(
                        f"Your application for the industrial attachment "
                        f"opportunity '{application.job.title}' was not "
                        f"successful. Thank you for applying."
                    ),
                    notification_type="APPLICATION_UPDATE",
                    audience="APPLICANT",
                )

            from accounts.audit import log_audit_event

            log_audit_event(
                request,
                "APPLICATION_DECISION_RECORDED",
                f"Application #{application.pk} for '{application.job.title}' "
                f"decided {new_status}",
                metadata={
                    "application_id": application.pk,
                    "applicant": application.user.username,
                    "previous_status": previous_status,
                    "new_status": new_status,
                },
            )

        return Response(
            {
                "detail": f"Application status updated to {new_status}.",
                "status": new_status,
            },
            status=status.HTTP_200_OK,
        )


# ---------------------------------------------------------------------------
# 1.4 Deployment & Exit Views
# ---------------------------------------------------------------------------

class DeploymentListCreateView(generics.ListCreateAPIView):
    serializer_class = DeploymentSerializer

    def get_permissions(self):
        if self.request.method == "POST":
            return [IsHRorAdmin()]
        return [IsDepartmentOrHRorAdmin()]

    def get_queryset(self):
        user = self.request.user
        qs = Deployment.objects.select_related(
            "application", "application__user", "application__job", "department", "created_by"
        ).order_by("-created_at")

        if is_director_role(user):
            # Strict scoping: a Director only ever sees deployments into their
            # own department. The scope comes from the authenticated user, and
            # a client-supplied `department` that is not their own is rejected
            # rather than trusted or silently ignored.
            director_dept = get_director_department(user)
            if director_dept is None:
                return Deployment.objects.none()

            requested_dept = self.request.query_params.get("department")
            if requested_dept and str(director_dept.id) != str(requested_dept):
                raise PermissionDenied(
                    "You can only view deployments for your own assigned department."
                )
            return qs.filter(department=director_dept)

        status_param = self.request.query_params.get("status")
        if status_param:
            qs = qs.filter(status=status_param)

        # Admin/HR remain institutional roles and may narrow to any department.
        department_id = self.request.query_params.get("department")
        if department_id and department_id != "ALL":
            qs = qs.filter(department_id=department_id)

        return qs

    def perform_create(self, serializer):
        application = serializer.validated_data["application"]

        if application.status != "SUCCESSFUL":
            from rest_framework.exceptions import ValidationError
            raise ValidationError(
                {"detail": "Only applicants whose application was SUCCESSFUL can be formally deployed to a department."}
            )

        start_date = serializer.validated_data.get("start_date")
        planned_end_date = serializer.validated_data.get("planned_end_date")
        if not planned_end_date and start_date:
            import datetime
            duration = getattr(application.job, "duration_weeks", 12) or 12
            planned_end_date = start_date + datetime.timedelta(weeks=duration)
            deployment = serializer.save(created_by=self.request.user, status="DEPLOYED", planned_end_date=planned_end_date)
        else:
            deployment = serializer.save(created_by=self.request.user, status="DEPLOYED")

        # Notify applicant of deployment
        Notification.objects.create(
            recipient=deployment.application.user,
            department=deployment.department,
            title="Formal Deployment Confirmed",
            message=(
                f"You have been deployed to the {deployment.department.name} department "
                f"starting from {deployment.start_date} to {deployment.planned_end_date}."
            ),
            notification_type="DEPLOYMENT_UPDATE",
            audience="APPLICANT",
        )


class DeploymentDetailView(generics.RetrieveUpdateAPIView):
    """Single deployment record.

    Read access for a Director is scoped server-side to their own department via
    ``get_queryset``, so requesting another department's deployment by ID returns
    404 rather than leaking the record.

    Writes stay with HR/Admin: ``department`` is a writable serializer field, so
    leaving update access with Directors would let them re-point their own
    deployment at a different department.
    """
    serializer_class = DeploymentSerializer

    def get_permissions(self):
        if self.request.method in permissions.SAFE_METHODS:
            return [IsDirectorOrHRorAdmin()]
        return [IsHRorAdmin()]

    def get_queryset(self):
        user = self.request.user
        qs = Deployment.objects.select_related(
            "application", "application__user", "application__job", "department", "created_by"
        )
        if is_director_role(user):
            director_dept = get_director_department(user)
            if director_dept is None:
                return Deployment.objects.none()
            return qs.filter(department=director_dept)
        return qs


class DeploymentExitView(APIView):
    permission_classes = [IsHRorAdmin]

    def post(self, request, pk):
        try:
            deployment = Deployment.objects.select_related(
                "application", "application__user", "application__job", "department"
            ).get(pk=pk)
        except Deployment.DoesNotExist:
            return Response({"detail": "Deployment not found."}, status=status.HTTP_404_NOT_FOUND)

        actual_end_date = request.data.get("actual_end_date") or timezone.now().date()
        deployment.actual_end_date = actual_end_date
        deployment.status = "EXITED"
        deployment.save()

        # Automatically initialize Dual Clearance in PENDING_DEPARTMENT state
        clearance, created = Clearance.objects.get_or_create(
            deployment=deployment,
            defaults={
                "application": deployment.application,
                "status": "PENDING_DEPARTMENT",
            }
        )
        if not created and clearance.status == "NOT_STARTED":
            clearance.status = "PENDING_DEPARTMENT"
            clearance.save()

        # Notify attachee and department
        Notification.objects.create(
            recipient=deployment.application.user,
            department=deployment.department,
            title="Attachment Completed - Clearance Unlocked",
            message=(
                f"Your attachment with {deployment.department.name} has concluded. "
                "Please submit your final report and begin your two-stage clearance process."
            ),
            notification_type="CLEARANCE_UPDATE",
            audience="APPLICANT",
        )

        return Response(
            {
                "detail": f"Deployment for {deployment.application.user.username} marked as EXITED. Dual clearance initialized.",
                "deployment": DeploymentSerializer(deployment).data,
            },
            status=status.HTTP_200_OK,
        )


class MyDeploymentView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        deployments = Deployment.objects.filter(
            application__user=request.user
        ).select_related("application", "application__job", "department", "created_by").order_by("-created_at")
        serializer = DeploymentSerializer(deployments, many=True)
        return Response({"deployments": serializer.data}, status=status.HTTP_200_OK)


# ---------------------------------------------------------------------------
# 1.5 Dual Clearance Views (Department Review -> HR Review)
# ---------------------------------------------------------------------------

class MyClearanceListView(generics.ListAPIView):
    serializer_class = ClearanceSerializer
    permission_classes = (permissions.IsAuthenticated,)

    def get_queryset(self):
        user = self.request.user
        return Clearance.objects.filter(
            Q(deployment__application__user=user) | Q(application__user=user)
        ).select_related(
            "deployment",
            "deployment__application",
            "deployment__application__job",
            "deployment__department",
            "application",
            "application__job",
            "department_cleared_by",
            "hr_cleared_by",
        ).order_by("-updated_at")


class ClearanceSubmitView(APIView):
    permission_classes = (permissions.IsAuthenticated,)

    def post(self, request, pk):
        user = request.user
        try:
            clearance = Clearance.objects.select_related(
                "deployment", "deployment__application", "deployment__department", "application", "application__job"
            ).get(
                Q(id=pk) & (Q(deployment__application__user=user) | Q(application__user=user))
            )
        except Clearance.DoesNotExist:
            return Response(
                {"detail": "Clearance record not found."},
                status=status.HTTP_404_NOT_FOUND,
            )

        if clearance.status == "CLEARED":
            return Response(
                {"detail": "Your clearance has already been approved and finalized."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # Gating: Final report submission is only enabled once Deployment.status == "EXITED"
        if clearance.deployment and clearance.deployment.status != "EXITED":
            return Response(
                {"detail": "Final report submission is only enabled once the industrial attachment deployment has reached completion and is marked EXITED."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        user_notes = request.data.get("user_notes", "")
        final_report_file = request.FILES.get("final_report_file")

        if final_report_file:
            ext = re.findall(r"\.[^.]+$", final_report_file.name)
            if not ext or ext[0].lower() not in [".pdf", ".docx", ".doc"]:
                return Response(
                    {"detail": "Only PDF and DOCX document formats are accepted for the final report."},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            if final_report_file.size > 10 * 1024 * 1024:
                return Response(
                    {"detail": "Final report file size must not exceed 10 MB."},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            clearance.final_report_file = final_report_file
            clearance.final_report_submitted_at = timezone.now()
        else:
            clearance.final_report_submitted_at = timezone.now()

        clearance.user_confirmed = True
        clearance.user_confirmed_at = timezone.now()
        clearance.user_notes = user_notes
        clearance.status = "PENDING_DEPARTMENT"
        clearance.save()

        dept = clearance.deployment.department if clearance.deployment else (clearance.application.job.department if clearance.application else None)
        Notification.objects.create(
            department=dept,
            title="Attachee Submitted Clearance Report",
            message=f"{user.first_name} {user.last_name} ({user.username}) submitted their report for Department clearance review.",
            notification_type="CLEARANCE_UPDATE",
        )

        return Response(
            {
                "detail": "Clearance documentation submitted successfully. Awaiting Department review.",
                "clearance": ClearanceSerializer(clearance).data,
            },
            status=status.HTTP_200_OK,
        )


class DepartmentClearanceQueueView(generics.ListAPIView):
    serializer_class = ClearanceSerializer
    permission_classes = [IsDirectorOrHRorAdmin]

    def get_queryset(self):
        user = self.request.user
        qs = Clearance.objects.select_related(
            "deployment", "deployment__application", "deployment__application__user",
            "deployment__department", "department_cleared_by", "hr_cleared_by"
        ).order_by("-updated_at")

        if user.role in ["DEPARTMENT_DIRECTOR", "DEPARTMENT"]:
            director_dept = get_director_department(user)
            if director_dept:
                qs = qs.filter(deployment__department=director_dept)
            else:
                return Clearance.objects.none()

        status_param = self.request.query_params.get("status")
        if status_param:
            qs = qs.filter(status=status_param)
        else:
            # Default to showing pending department review
            qs = qs.filter(status__in=["PENDING_DEPARTMENT", "REJECTED"])

        return qs


class DepartmentClearanceReviewView(APIView):
    permission_classes = [IsDirectorOrHRorAdmin]

    def post(self, request, pk):
        try:
            clearance = Clearance.objects.select_related(
                "deployment", "deployment__application", "deployment__department"
            ).get(pk=pk)
        except Clearance.DoesNotExist:
            return Response({"detail": "Clearance record not found."}, status=status.HTTP_404_NOT_FOUND)

        # If user is a Department Director, verify the clearance belongs to them.
        if is_director_role(request.user):
            director_dept = get_director_department(request.user)
            if not director_dept:
                return Response(
                    {"detail": "You do not have an assigned department to clear attachees for."},
                    status=status.HTTP_403_FORBIDDEN,
                )
            # A clearance may be raised from a deployment or, for legacy records,
            # straight from an application. Both paths must resolve to the
            # Director's own department before any sign-off is accepted.
            if clearance.deployment:
                clearance_dept = clearance.deployment.department
            elif clearance.application and clearance.application.job:
                clearance_dept = clearance.application.job.department
            else:
                clearance_dept = None

            if clearance_dept is None or clearance_dept != director_dept:
                return Response(
                    {"detail": "You do not have permission to clear attachees outside your assigned department."},
                    status=status.HTTP_403_FORBIDDEN,
                )

        decision = request.data.get("decision", "").upper()
        notes = request.data.get("department_notes", "")

        if decision == "APPROVE":
            clearance.department_cleared = True
            clearance.department_cleared_by = request.user
            clearance.department_cleared_at = timezone.now()
            clearance.department_notes = notes
            clearance.status = "PENDING_HR"
            clearance.save()

            # Notify HR of pending final review
            Notification.objects.create(
                department=clearance.deployment.department if clearance.deployment else None,
                title="Department Clearance Approved",
                message=f"Department sign-off granted for {clearance.deployment.application.user.username}. Ready for final HR clearance.",
                notification_type="CLEARANCE_UPDATE",
            )
            return Response(
                {"detail": "Department clearance approved. Record forwarded to HR queue.", "status": "PENDING_HR"},
                status=status.HTTP_200_OK,
            )

        elif decision == "REJECT":
            clearance.department_cleared = False
            clearance.department_notes = notes
            clearance.status = "REJECTED"
            clearance.save()

            if clearance.deployment:
                Notification.objects.create(
                    recipient=clearance.deployment.application.user,
                    title="Department Clearance Revision Requested",
                    message=f"Your department has requested revisions: {notes}",
                    notification_type="CLEARANCE_UPDATE",
                    audience="APPLICANT",
                )
            return Response(
                {"detail": "Department clearance rejected / revisions requested.", "status": "REJECTED"},
                status=status.HTTP_200_OK,
            )

        return Response({"detail": "Invalid decision. Must be 'APPROVE' or 'REJECT'."}, status=status.HTTP_400_BAD_REQUEST)


class HRClearanceQueueView(generics.ListAPIView):
    serializer_class = ClearanceSerializer
    permission_classes = [IsHRorAdmin]

    def get_queryset(self):
        qs = Clearance.objects.select_related(
            "deployment", "deployment__application", "deployment__application__user",
            "deployment__department", "department_cleared_by", "hr_cleared_by"
        ).order_by("-updated_at")

        status_param = self.request.query_params.get("status")
        if status_param:
            qs = qs.filter(status=status_param)
        else:
            qs = qs.filter(status="PENDING_HR")

        department_id = self.request.query_params.get("department")
        if department_id:
            qs = qs.filter(deployment__department_id=department_id)

        return qs


class HRClearanceReviewView(APIView):
    permission_classes = [IsHRorAdmin]

    def post(self, request, pk):
        try:
            clearance = Clearance.objects.select_related(
                "deployment", "deployment__application", "deployment__department"
            ).get(pk=pk)
        except Clearance.DoesNotExist:
            return Response({"detail": "Clearance record not found."}, status=status.HTTP_404_NOT_FOUND)

        decision = request.data.get("decision", "").upper()
        notes = request.data.get("hr_notes", "")

        if decision == "APPROVE":
            # Server-side sequence enforcement: Department clearance MUST be granted first
            if not clearance.department_cleared:
                return Response(
                    {"detail": "Cannot grant HR clearance. Department-level sign-off must be completed first."},
                    status=status.HTTP_400_BAD_REQUEST,
                )

            clearance.hr_cleared = True
            clearance.hr_cleared_by = request.user
            clearance.hr_cleared_at = timezone.now()
            clearance.hr_notes = notes
            clearance.status = "CLEARED"
            clearance.save()

            if clearance.deployment:
                Notification.objects.create(
                    recipient=clearance.deployment.application.user,
                    department=clearance.deployment.department,
                    title="Official Clearance Granted!",
                    message="Congratulations! Your institutional clearance has been fully approved by HR.",
                    notification_type="CLEARANCE_UPDATE",
                    audience="APPLICANT",
                )
            return Response(
                {"detail": "HR Clearance approved. Full clearance finalized.", "status": "CLEARED"},
                status=status.HTTP_200_OK,
            )

        elif decision == "REJECT":
            clearance.hr_cleared = False
            clearance.hr_notes = notes
            clearance.status = "REJECTED"
            clearance.save()

            if clearance.deployment:
                Notification.objects.create(
                    recipient=clearance.deployment.application.user,
                    title="HR Clearance Revision Requested",
                    message=f"HR has requested revisions: {notes}",
                    notification_type="CLEARANCE_UPDATE",
                    audience="APPLICANT",
                )
            return Response(
                {"detail": "HR clearance rejected / revisions requested.", "status": "REJECTED"},
                status=status.HTTP_200_OK,
            )

        return Response({"detail": "Invalid decision. Must be 'APPROVE' or 'REJECT'."}, status=status.HTTP_400_BAD_REQUEST)


# ---------------------------------------------------------------------------
# 1.2 & 1.6 Recommendation Letter Templates & Letter Generation Views
# ---------------------------------------------------------------------------

class RecommendationLetterTemplateListCreateView(generics.ListCreateAPIView):
    serializer_class = RecommendationLetterTemplateSerializer
    permission_classes = [IsHRorAdmin]

    def get_queryset(self):
        return RecommendationLetterTemplate.objects.all().order_by("-is_default", "-created_at")

    def perform_create(self, serializer):
        is_default = serializer.validated_data.get("is_default", False)
        if is_default:
            RecommendationLetterTemplate.objects.filter(is_default=True).update(is_default=False)
        serializer.save(created_by=self.request.user)


class RecommendationLetterTemplateDetailView(generics.RetrieveUpdateDestroyAPIView):
    queryset = RecommendationLetterTemplate.objects.all()
    serializer_class = RecommendationLetterTemplateSerializer
    permission_classes = [IsHRorAdmin]

    def perform_update(self, serializer):
        is_default = serializer.validated_data.get("is_default", False)
        if is_default:
            RecommendationLetterTemplate.objects.filter(is_default=True).exclude(pk=self.get_object().pk).update(is_default=False)
        serializer.save()


class RecommendationLetterGenerateView(APIView):
    permission_classes = [IsHRorAdmin]

    def post(self, request, deployment_id):
        try:
            deployment = Deployment.objects.select_related(
                "application", "application__user", "application__user__profile",
                "application__job", "department", "clearance"
            ).get(pk=deployment_id)
        except Deployment.DoesNotExist:
            return Response({"detail": "Deployment not found."}, status=status.HTTP_404_NOT_FOUND)

        # Server-side validation: must have status CLEARED
        clearance = getattr(deployment, "clearance", None)
        if not clearance or clearance.status != "CLEARED":
            status_text = clearance.status if clearance else "NOT_STARTED"
            return Response(
                {"detail": f"Institutional clearance must be CLEARED before generating a recommendation letter (current status: {status_text})."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # Server-side validation: final report must be submitted
        if not clearance.final_report_submitted_at and not clearance.final_report_file:
            return Response(
                {"detail": "Attachee final report has not been submitted. A submitted final report is required before generating a recommendation letter."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # Retrieve template
        template_id = request.data.get("template_id")
        if template_id:
            try:
                template = RecommendationLetterTemplate.objects.get(pk=template_id)
            except RecommendationLetterTemplate.DoesNotExist:
                template = RecommendationLetterTemplate.objects.filter(is_default=True).first()
        else:
            template = RecommendationLetterTemplate.objects.filter(is_default=True).first()

        if not template:
            # Fallback standard template body if none created yet
            template_body = (
                "STATE DEPARTMENT FOR PETROLEUM\n"
                "MINISTRY OF ENERGY AND PETROLEUM\n\n"
                "Date: {{issue_date}}\n\n"
                "TO WHOM IT MAY CONCERN\n\n"
                "RE: RECOMMENDATION FOR {{full_name}}\n\n"
                "This is to certify that {{full_name}} from {{institution_name}} successfully undertook and completed "
                "an Industrial Attachment programme in the {{department}} as {{role_or_position}} "
                "from {{start_date}} to {{end_date}}.\n\n"
                "During their tenure, {{full_name}} demonstrated outstanding diligence, technical proficiency, "
                "and exemplary professional conduct. All institutional clearance requirements and departmental obligations "
                "have been fully discharged.\n\n"
                "We highly recommend {{full_name}} for future academic and professional engagements.\n\n"
                "Sincerely,\n\n"
                "{{hr_name}}\n"
                "Directorate of Human Resource Management\n"
                "State Department for Petroleum"
            )
        else:
            template_body = template.body

        # Data mapping for token substitution
        user = deployment.application.user
        profile = getattr(user, "profile", None)
        edu = user.education.first()
        institution_name = (profile.institution_name if profile and profile.institution_name else None) or (edu.institution_name if edu else None) or "University / Tertiary Institution"
        start_date_str = deployment.start_date.strftime("%B %d, %Y") if deployment.start_date else "N/A"
        end_date_str = (deployment.actual_end_date or deployment.planned_end_date).strftime("%B %d, %Y")
        issue_date_str = timezone.now().strftime("%B %d, %Y")
        hr_name_str = f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username

        tokens = {
            "{{full_name}}": f"{user.first_name} {user.last_name}".strip() or user.username,
            "{{department}}": deployment.department.name,
            "{{role_or_position}}": deployment.application.job.title,
            "{{start_date}}": start_date_str,
            "{{end_date}}": end_date_str,
            "{{issue_date}}": issue_date_str,
            "{{hr_name}}": hr_name_str,
            "{{institution_name}}": institution_name,
        }

        rendered = template_body
        for token, val in tokens.items():
            rendered = rendered.replace(token, str(val))

        letter = RecommendationLetter.objects.create(
            deployment=deployment,
            template_used=template if template else None,
            rendered_content=rendered,
            generated_by=request.user,
        )

        return Response(
            {
                "detail": "Recommendation letter generated successfully.",
                "letter": RecommendationLetterSerializer(letter).data,
            },
            status=status.HTTP_201_CREATED,
        )


class RecommendationLetterListView(generics.ListAPIView):
    serializer_class = RecommendationLetterSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        base = RecommendationLetter.objects.all().select_related(
            "deployment", "deployment__application", "deployment__application__user",
            "deployment__department", "template_used", "generated_by"
        )
        if user.role in ["ADMIN", "HR"]:
            return base.order_by("-generated_at")
        # Directors were missing here even though the detail view and the
        # download view both scope them to their department. A Director could
        # open a letter by its id but the list they were expected to browse came
        # back empty, so the department's letters were unreachable in practice.
        if user.role in ["DEPARTMENT_DIRECTOR", "DEPARTMENT"]:
            dept = get_director_department(user)
            if not dept:
                return RecommendationLetter.objects.none()
            return base.filter(deployment__department=dept).order_by("-generated_at")
        return base.filter(
            deployment__application__user=user
        ).order_by("-generated_at")


class RecommendationLetterDetailView(generics.RetrieveAPIView):
    serializer_class = RecommendationLetterSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        if user.role in ["ADMIN", "HR"]:
            return RecommendationLetter.objects.all().select_related(
                "deployment", "deployment__application", "deployment__application__user",
                "deployment__department", "template_used", "generated_by"
            )
        elif user.role in ["DEPARTMENT_DIRECTOR", "DEPARTMENT"]:
            dept = get_director_department(user)
            return RecommendationLetter.objects.filter(
                deployment__department=dept
            ).select_related(
                "deployment", "deployment__application", "deployment__application__user",
                "deployment__department", "template_used", "generated_by"
            ) if dept else RecommendationLetter.objects.none()
        return RecommendationLetter.objects.filter(
            deployment__application__user=user
        ).select_related(
            "deployment", "deployment__application", "deployment__application__user",
            "deployment__department", "template_used", "generated_by"
        )


# ---------------------------------------------------------------------------
# Recommendation Letter DOCX Download
# ---------------------------------------------------------------------------

class RecommendationLetterDownloadView(APIView):
    """
    Download a recommendation letter as a professionally formatted DOCX document.
    GET /api/jobs/recommendations/<pk>/download/
    """
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request, pk):
        from docx import Document as DocxDocument
        from docx.shared import Inches, Pt, RGBColor, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.section import WD_ORIENT
        import io

        user = request.user
        try:
            letter = RecommendationLetter.objects.select_related(
                "deployment", "deployment__application", "deployment__application__user",
                "deployment__department", "template_used", "generated_by"
            ).get(pk=pk)
        except RecommendationLetter.DoesNotExist:
            return Response({"detail": "Recommendation letter not found."}, status=status.HTTP_404_NOT_FOUND)

        # Permission check
        if user.role not in ["ADMIN", "HR"]:
            if is_director_role(user):
                dept = get_director_department(user)
                if not dept or letter.deployment.department != dept:
                    return Response({"detail": "Permission denied."}, status=status.HTTP_403_FORBIDDEN)
            elif letter.deployment.application.user != user:
                return Response({"detail": "Permission denied."}, status=status.HTTP_403_FORBIDDEN)

        # Build DOCX
        doc = DocxDocument()

        # Page setup
        section = doc.sections[0]
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

        # ── Republic of Kenya Coat of Arms header ──
        header_line1 = doc.add_paragraph()
        header_line1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run1 = header_line1.add_run("REPUBLIC OF KENYA")
        run1.font.size = Pt(16)
        run1.font.bold = True
        run1.font.color.rgb = RGBColor(0, 51, 102)

        header_line2 = doc.add_paragraph()
        header_line2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run2 = header_line2.add_run("STATE DEPARTMENT FOR PETROLEUM")
        run2.font.size = Pt(13)
        run2.font.bold = True
        run2.font.color.rgb = RGBColor(0, 51, 102)

        header_line3 = doc.add_paragraph()
        header_line3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run3 = header_line3.add_run("MINISTRY OF ENERGY AND PETROLEUM")
        run3.font.size = Pt(11)
        run3.font.bold = True
        run3.font.color.rgb = RGBColor(80, 80, 80)

        header_line4 = doc.add_paragraph()
        header_line4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run4 = header_line4.add_run("Industrial Attachment System")
        run4.font.size = Pt(10)
        run4.font.italic = True
        run4.font.color.rgb = RGBColor(100, 100, 100)

        # Horizontal divider
        divider = doc.add_paragraph()
        divider.alignment = WD_ALIGN_PARAGRAPH.CENTER
        div_run = divider.add_run("─" * 60)
        div_run.font.size = Pt(8)
        div_run.font.color.rgb = RGBColor(180, 180, 180)

        # ── Letter body ──
        # Split rendered_content by lines and add as paragraphs
        content_lines = letter.rendered_content.split("\n")
        for line in content_lines:
            stripped = line.strip()
            p = doc.add_paragraph()

            if stripped.startswith("RE:") or stripped.startswith("TO WHOM"):
                # Heading-style lines
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(stripped)
                run.font.bold = True
                run.font.size = Pt(12)
                run.font.color.rgb = RGBColor(0, 51, 102)
            elif stripped.startswith("Date:"):
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(stripped)
                run.font.size = Pt(11)
                run.font.italic = True
            elif stripped.startswith("STATE DEPARTMENT") or stripped.startswith("MINISTRY OF"):
                # Already in header, skip duplicates
                continue
            elif stripped.startswith("Sincerely") or stripped.startswith("Yours"):
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(stripped)
                run.font.size = Pt(11)
                run.font.italic = True
            elif stripped == "":
                # Empty line spacing
                p.space_after = Pt(6)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                run = p.add_run(stripped)
                run.font.size = Pt(11)
                run.font.name = "Calibri"

        # ── Footer divider ──
        footer_div = doc.add_paragraph()
        footer_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fd_run = footer_div.add_run("─" * 60)
        fd_run.font.size = Pt(8)
        fd_run.font.color.rgb = RGBColor(180, 180, 180)

        # ── Footer ──
        footer_p = doc.add_paragraph()
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fr = footer_p.add_run(
            f"Generated on {letter.generated_at.strftime('%B %d, %Y at %I:%M %p')} | "
            f"Industrial Attachment System — State Department for Petroleum"
        )
        fr.font.size = Pt(8)
        fr.font.italic = True
        fr.font.color.rgb = RGBColor(130, 130, 130)

        ref_p = doc.add_paragraph()
        ref_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ref_run = ref_p.add_run(f"Reference No: IAS/REC/{letter.pk:05d}/{letter.generated_at.strftime('%Y')}")
        ref_run.font.size = Pt(8)
        ref_run.font.bold = True
        ref_run.font.color.rgb = RGBColor(0, 51, 102)

        # Generate file
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        applicant_name = f"{letter.deployment.application.user.first_name}_{letter.deployment.application.user.last_name}".strip("_") or "Attachee"
        filename = f"Recommendation_Letter_{applicant_name}_{letter.generated_at.strftime('%Y%m%d')}.docx"

        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response


# ---------------------------------------------------------------------------
# Notifications Views
# ---------------------------------------------------------------------------

def notifications_visible_to(user):
    """
    The notifications a given user is allowed to see.

    Defined once and reused by the list, mark-read and mark-all-read views.
    They previously each rebuilt this rule inline, which is how the three drifted
    apart: applicant-private notices were excluded from none of them, so HR and
    Admin received every applicant's outcome and the Director of the vacancy's
    department received them too.

    The rule:
      * Applicants see only notices addressed to them, which is where their
        application outcome lands.
      * Admins and HR see the staff feed, which is the whole system minus
        applicant-private notices. They keep every operational notice.
      * Directors see their own department's staff notices plus anything addressed
        to them personally.
    """
    qs = Notification.objects.select_related("department", "recipient")

    if user.role in ["ADMIN", "HR"]:
        return qs.exclude(audience="APPLICANT")

    if is_director_role(user):
        dept = get_director_department(user)
        if dept:
            return qs.filter(Q(department=dept) | Q(recipient=user)).exclude(audience="APPLICANT")

    return qs.filter(recipient=user)


class NotificationListView(generics.ListAPIView):
    serializer_class = NotificationSerializer
    permission_classes = (permissions.IsAuthenticated,)

    def get_queryset(self):
        return notifications_visible_to(self.request.user).order_by("-created_at")[:50]


class NotificationMarkReadView(APIView):
    permission_classes = (permissions.IsAuthenticated,)

    def post(self, request, pk):
        return self.patch(request, pk)

    def patch(self, request, pk):
        user = request.user
        try:
            # Scoped through the same helper, so a notification the user cannot
            # see in the list cannot be marked read by guessing its id.
            notification = notifications_visible_to(user).get(pk=pk)

            notification.is_read = True
            notification.save()
            return Response({"detail": "Notification marked as read."}, status=status.HTTP_200_OK)
        except Notification.DoesNotExist:
            return Response({"detail": "Notification not found."}, status=status.HTTP_404_NOT_FOUND)


class NotificationMarkAllReadView(APIView):
    permission_classes = (permissions.IsAuthenticated,)

    def post(self, request):
        # Same scope, so "mark all read" never silently consumes an applicant's
        # private notices on the staff side.
        notifications_visible_to(request.user).filter(is_read=False).update(is_read=True)
        return Response({"detail": "All notifications marked as read."}, status=status.HTTP_200_OK)


# ---------------------------------------------------------------------------
# Reporting & CSV Export Views
# ---------------------------------------------------------------------------
#
# The JSON report endpoints and the CSV export endpoint are two presentations of
# the same report-generation logic.  They therefore share the queryset builders
# and row serializers defined here, so a department-scope fix can never be
# applied to one presentation and forgotten on the other.

# Query parameters the attachee report accepts in addition to the department.
ATTACHEE_REPORT_STATUS_PARAM = "status"
ATTACHEE_REPORT_START_PARAM = "start_date"
ATTACHEE_REPORT_END_PARAM = "end_date"


# The single definition of the attachee report's columns. The JSON report, the
# CSV export and the print view all read this list, so a field can never be
# present on screen but missing from the export (or vice versa) — the class of
# bug that makes an official export unusable as a record.
#
# Each entry is (json_key, csv_header). Keys are also what the frontend maps
# over, so the column set is identical on both sides of the API.
ATTACHEE_REPORT_COLUMNS = [
    ("application_id", "Application ID"),
    ("applicant_name", "Attachee Name"),
    ("username", "Username"),
    ("applicant_email", "Email"),
    ("phone_number", "Phone"),
    ("id_number", "National ID"),
    ("county_of_residence", "County"),
    ("institution_name", "Institution"),
    ("qualification", "Qualification"),
    ("field_of_study", "Field of Study"),
    ("opportunity_type", "Opportunity Type"),
    ("job_title", "Vacancy Title"),
    ("department_name", "Department"),
    ("job_location", "Attachment Location"),
    ("duration_weeks", "Duration (Weeks)"),
    ("application_date", "Application Date"),
    # "hired_date" is deliberately not exported as its own column. It carries
    # the same value as "application_date" and was labelled "Placement Date",
    # so the export printed the application date twice under a name that
    # described neither it nor any real placement event.
    ("deployment_start_date", "Deployment Start Date"),
    ("deployment_end_date", "Deployment End Date"),
    ("ats_score", "ATS Score (%)"),
    ("clearance_status", "Clearance Status"),
    ("department_cleared", "Department Cleared"),
    ("department_cleared_at", "Department Clearance Date"),
    ("hr_cleared", "HR Approved"),
    ("hr_cleared_by", "HR Cleared By"),
    ("hr_cleared_at", "HR Clearance Date"),
    ("final_report_submitted", "Final Report Submitted"),
    ("recommendation_letter", "Recommendation Letter"),
    ("is_completed", "Attachment Completed"),
    ("completion_date", "Completion Date"),
    ("is_archived", "Archived"),
]

FILL_RATE_REPORT_COLUMNS = [
    ("job_id", "Vacancy ID"),
    ("title", "Vacancy Title"),
    ("department_name", "Department"),
    ("job_type", "Opportunity Type"),
    ("opportunity_type", "Opportunity Label"),
    ("location", "Location"),
    ("slots_required", "Slots Required"),
    ("slots_filled", "Slots Filled"),
    ("slots_remaining", "Slots Remaining"),
    ("fill_rate", "Fill Rate (%)"),
    ("is_full", "Capacity Full"),
    ("is_archived", "Archived"),
    ("duration_weeks", "Duration (Weeks)"),
    ("created_at", "Posted On"),
    ("deadline", "Deadline"),
]

#: Rendered instead of an empty cell. An export that silently leaves a cell
#: blank reads as "we do not know this", whereas an explicit "Not recorded"
#: distinguishes a genuinely missing value from a dropped column.
NOT_RECORDED = "Not recorded"


def complete_cell(value, blank_placeholder=NOT_RECORDED):
    """
    Normalises a single export cell so no column is ever silently blank.

    ``None``, an empty string and whitespace all become the same explicit
    placeholder, and every other value is rendered as text. Booleans become
    Yes/No so the export is readable rather than a column of True/False.
    """
    if value is None:
        return blank_placeholder
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, str):
        return value if value.strip() else blank_placeholder
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def build_export_row(record, columns, blank_placeholder=NOT_RECORDED):
    """Render one record against a column list, with no cell left empty."""
    return [complete_cell(record.get(key), blank_placeholder) for key, _ in columns]


def column_headers(columns):
    return [header for _, header in columns]


def build_attachees_report_queryset(request):
    """Successfully selected attachees for the requested scope, with clearance state attached.

    Shared by ``AttacheesReportView`` and ``ExportCSVReportView``.
    """
    qs = Application.objects.filter(status="SUCCESSFUL").select_related(
        "user", "user__profile", "job", "job__department", "clearance", "deployment"
    ).order_by("-applied_at")

    qs = scope_queryset_to_request_department(request, qs, "job__department")

    clearance_status = request.query_params.get(ATTACHEE_REPORT_STATUS_PARAM)
    if clearance_status:
        qs = qs.filter(clearance__status=clearance_status)

    start_date = request.query_params.get(ATTACHEE_REPORT_START_PARAM)
    if start_date:
        qs = qs.filter(applied_at__date__gte=start_date)

    end_date = request.query_params.get(ATTACHEE_REPORT_END_PARAM)
    if end_date:
        qs = qs.filter(applied_at__date__lte=end_date)

    return qs


def build_fill_rates_report_queryset(request):
    """Vacancies for the requested scope, with derived slot fill rates.

    Shared by ``FillRatesReportView`` and ``ExportCSVReportView``.
    """
    qs = Job.objects.select_related("department").order_by("-created_at")
    qs = scope_queryset_to_request_department(request, qs, "department")

    job_type = request.query_params.get("job_type")
    if job_type:
        qs = qs.filter(job_type=job_type)

    return qs


def serialize_attachee_report_row(app):
    """
    One attachee row, as returned by the JSON report and the CSV export.

    Every key in ``ATTACHEE_REPORT_COLUMNS`` is present here, so the on-screen
    report, the CSV export and the print view can never disagree about which
    fields exist. A genuinely unknown value is returned as ``None`` (and
    rendered as an explicit "Not recorded" cell downstream) rather than being
    omitted, so a column is never silently missing from a record.
    """
    from .attachment_records import is_completed_attachment

    user = app.user
    profile = getattr(user, "profile", None)
    job = app.job
    department = job.department if job else None
    clearance = getattr(app, "clearance", None)
    deployment = getattr(app, "deployment", None)

    duration_weeks = job.duration_weeks if job else None
    if duration_weeks is None and deployment and deployment.start_date and deployment.actual_end_date:
        duration_weeks = round(
            (deployment.actual_end_date - deployment.start_date).days / 7
        ) or None

    # Evaluated once: it touches the deployment and clearance relations, and the
    # completion date below is only meaningful when it is True.
    completed = is_completed_attachment(app)

    return {
        "application_id": app.id,
        "applicant_name": f"{user.first_name} {user.last_name}".strip() or user.username,
        "username": user.username,
        "applicant_email": user.email or None,
        "phone_number": (profile.phone_number if profile else None) or None,
        "id_number": (profile.id_number if profile else None) or None,
        "county_of_residence": (profile.county_of_residence if profile else None) or None,
        "institution_name": (profile.institution_name if profile else None) or None,
        "qualification": (profile.qualification if profile else None) or None,
        "field_of_study": (profile.field_of_study if profile else None) or None,
        "opportunity_type": job.get_job_type_display() if job else None,
        "job_title": job.title if job else None,
        "department_name": department.name if department else "General",
        "job_location": (job.location if job else None) or None,
        "duration_weeks": duration_weeks,
        "application_date": app.applied_at.date() if app.applied_at else None,
        # Kept for backwards compatibility with the original report payload.
        "hired_date": app.applied_at.strftime("%Y-%m-%d") if app.applied_at else None,
        "deployment_start_date": deployment.start_date if deployment else None,
        "deployment_end_date": deployment.actual_end_date if deployment else None,
        "ats_score": float(app.ats_score) if app.ats_score is not None else None,
        "clearance_status": clearance.status if clearance else "NOT_STARTED",
        "department_cleared": clearance.department_cleared if clearance else False,
        "department_cleared_at": clearance.department_cleared_at if clearance else None,
        "hr_cleared": clearance.hr_cleared if clearance else False,
        "hr_cleared_by": clearance.hr_cleared_by.username if (clearance and clearance.hr_cleared_by) else None,
        "hr_cleared_at": clearance.hr_cleared_at if clearance else None,
        "final_report_submitted": bool(clearance.final_report_file) if clearance else False,
        "recommendation_letter": bool(getattr(app, "_has_recommendation_letter", False)),
        "is_completed": completed,
        "completion_date": (
            clearance.hr_cleared_at.date()
            if (completed and clearance and clearance.hr_cleared_at)
            else None
        ),
        "is_archived": bool(app.is_archived),
    }


def serialize_fill_rate_row(job):
    """
    One vacancy fill-rate row, as returned by the JSON report and CSV export.

    Keys match ``FILL_RATE_REPORT_COLUMNS`` exactly, for the same reason as the
    attachee report: the export must carry every field the report shows.
    """
    slots_required = job.slots_required or 0
    slots_filled = job.slots_filled
    fill_rate = round((slots_filled / slots_required) * 100, 1) if slots_required > 0 else 0
    return {
        "job_id": job.id,
        "title": job.title,
        "department_name": job.department.name if job.department else "General",
        "job_type": job.job_type,
        "opportunity_type": job.get_job_type_display(),
        "location": job.location or None,
        "slots_required": slots_required,
        "slots_filled": slots_filled,
        "slots_remaining": max(slots_required - slots_filled, 0),
        "fill_rate": f"{fill_rate}%" if slots_required > 0 else None,
        "is_full": job.is_full,
        "is_archived": job.is_archived,
        "duration_weeks": job.duration_weeks,
        "created_at": job.created_at if job.created_at else None,
        "deadline": job.deadline if job.deadline else None,
    }


class AttacheesReportView(APIView):
    """Attachee & clearance report.

    Admin/HR: system-wide, or scoped to any department via ``?department=``.
    Directors: always and unconditionally scoped to their own department, derived
    from the authenticated user.  A Director requesting another department is
    rejected with 403 rather than having the client-supplied value trusted.
    """

    def get_permissions(self):
        return [IsDirectorOrHRorAdmin()]

    def get(self, request):
        results = [
            serialize_attachee_report_row(app)
            for app in build_attachees_report_queryset(request)
        ]
        return Response(
            {
                "attachees": results,
                "total_count": len(results),
                # The same column definition the CSV export is generated from, so
                # the client renders and exports an identical field set.
                "columns": [
                    {"key": key, "header": header}
                    for key, header in ATTACHEE_REPORT_COLUMNS
                ],
            },
            status=status.HTTP_200_OK,
        )


class FillRatesReportView(APIView):
    """Vacancy slot fill-rate report.

    Admin/HR: system-wide, or scoped to any department via ``?department=``.
    Directors: always and unconditionally scoped to their own department.
    """

    def get_permissions(self):
        return [IsDirectorOrHRorAdmin()]

    def get(self, request):
        results = [
            serialize_fill_rate_row(job)
            for job in build_fill_rates_report_queryset(request)
        ]
        return Response(
            {
                "vacancies": results,
                "total_count": len(results),
                "columns": [
                    {"key": key, "header": header}
                    for key, header in FILL_RATE_REPORT_COLUMNS
                ],
            },
            status=status.HTTP_200_OK,
        )


def write_fill_rates_csv(writer, qs):
    """
    Emit the fill-rate report body as CSV rows; returns the row count.

    Driven by ``FILL_RATE_REPORT_COLUMNS`` and ``build_export_row`` so the CSV
    header, the CSV body and the on-screen report always describe the same set
    of fields, with no cell left blank.
    """
    writer.writerow(column_headers(FILL_RATE_REPORT_COLUMNS))
    row_count = 0
    for job in qs:
        writer.writerow(
            build_export_row(serialize_fill_rate_row(job), FILL_RATE_REPORT_COLUMNS)
        )
        row_count += 1
    return row_count


def write_attachees_csv(writer, qs):
    """
    Emit the attachee report body as CSV rows; returns the row count.

    Driven by ``ATTACHEE_REPORT_COLUMNS`` and ``build_export_row`` so every
    field the report shows is present in the export, and so an unknown value is
    written as an explicit "Not recorded" instead of a blank cell that looks
    like a missing column.
    """
    writer.writerow(column_headers(ATTACHEE_REPORT_COLUMNS))
    row_count = 0
    for app in qs:
        writer.writerow(
            build_export_row(serialize_attachee_report_row(app), ATTACHEE_REPORT_COLUMNS)
        )
        row_count += 1
    return row_count


class ExportCSVReportView(APIView):
    """CSV export of the attachee and fill-rate reports.

    Uses the exact same queryset builders as the JSON report endpoints, so the
    CSV can never widen the department scope that the on-screen report enforces.
    """

    def get_permissions(self):
        return [IsDirectorOrHRorAdmin()]

    def get(self, request):
        report_type = request.query_params.get("report_type", "attachees")
        department = resolve_requested_department(request)

        timestamp = timezone.now().strftime("%Y%m%d_%H%M")
        now_display = timezone.now().strftime("%d %B %Y at %H:%M")
        response = HttpResponse(content_type="text/csv")
        response["Content-Disposition"] = f'attachment; filename="{report_type}_report_{timestamp}.csv"'
        writer = csv.writer(response)

        # Professional header block
        report_title = "Vacancy Fill Rates Report" if report_type == "fill_rates" else "Attachees & Clearance Report"
        report_scope = f" — {department.name}" if department else ""
        writer.writerow([""])
        writer.writerow(["REPUBLIC OF KENYA"])
        writer.writerow(["STATE DEPARTMENT FOR PETROLEUM"])
        writer.writerow(["INDUSTRIAL ATTACHMENT SYSTEM"])
        writer.writerow([""])
        writer.writerow([f"Report:  {report_title}{report_scope}"])
        writer.writerow([f"Generated:  {now_display}"])
        writer.writerow([""])
        writer.writerow(["=" * 80])
        writer.writerow([""])

        if report_type == "fill_rates":
            row_count = write_fill_rates_csv(writer, build_fill_rates_report_queryset(request))
        else:
            row_count = write_attachees_csv(writer, build_attachees_report_queryset(request))

        # Professional footer block
        writer.writerow([""])
        writer.writerow(["-" * 80])
        writer.writerow([f"End of Report — {report_title} — {row_count} records"])
        writer.writerow([f"Total Records:  {row_count}"])
        writer.writerow(["State Department for Petroleum — Industrial Attachment System"])

        return response


def attachee_display_name(user):
    """The attachee's name as a person, falling back to their username."""
    if user is None:
        return "Attachee"
    full_name = f"{user.first_name} {user.last_name}".strip()
    return full_name or user.username


def clearance_attachee_name(clearance):
    """
    The attachee a clearance belongs to.

    A Clearance may hang off a deployment, or -- for legacy rows that predate
    deployments -- straight off an application. Both are handled so this never
    returns a name for the wrong person, and never raises on a missing relation.
    """
    application = clearance.deployment.application if clearance.deployment else clearance.application
    return attachee_display_name(application.user if application else None)


def job_capacity_totals(job_qs):
    """
    Total required slots and total filled slots for a set of vacancies.

    ``Job.slots_filled`` is a property that issues a COUNT per job, so summing
    it across a queryset meant one extra query per vacancy. The filled total is
    computed here with a single grouped query instead.

    Filled means "selected successfully", matching ``Job.slots_filled``, so the
    capacity figures here and the per-vacancy fill rates in the reports cannot
    disagree.
    """
    required = sum(job_qs.values_list("slots_required", flat=True)) or 0
    filled = (
        Application.objects.filter(job__in=job_qs, status="SUCCESSFUL")
        .values("job_id")
        .annotate(successful=Count("id"))
    )
    filled_total = sum(row["successful"] for row in filled)
    return required, filled_total


def calculate_statistical_metrics(scores_list):
    """
    Computes rigorous actuarial & descriptive statistical metrics for a list of numerical scores.
    Handles population/sample variance, quartiles, interquartile range (IQR), and tier distribution.

    Quality tier boundaries and the qualification benchmark are NOT hardcoded here: they are
    read from the singleton AtsScoringConfiguration so that Admin-configured values in
    System Settings & Keys are the single source of truth for tiering.

    Every figure is computed from the unrounded data and rounded exactly once at
    the end. The previous implementation squared the *rounded* standard deviation
    to obtain the variance, which could be wrong by a couple of hundredths
    before rounding even took effect (a stdev of 12.3456 reported 12.35, whose
    square is 152.52 against a true variance of 152.41).
    """
    ats_config = AtsScoringConfiguration.get_config()
    clean_scores = ats_config.sanitize_scores(scores_list)
    n = len(clean_scores)

    if n == 0:
        return {
            "count": 0,
            "mean": 0.0,
            "median": 0.0,
            "std_dev": 0.0,
            "variance": 0.0,
            "min_score": 0.0,
            "max_score": 0.0,
            "q1": 0.0,
            "q3": 0.0,
            "iqr": 0.0,
            "qualification_rate": 0.0,
            "distribution_buckets": ats_config.empty_distribution_buckets(),
        }

    sorted_scores = sorted(clean_scores)
    min_raw = min(clean_scores)
    max_raw = max(clean_scores)

    # Sample variance for n > 1, and 0 for a single observation (where the
    # sample variance is undefined). Stated once here and used for both the
    # variance and the standard deviation so the two cannot disagree.
    variance_raw = statistics.variance(clean_scores) if n > 1 else 0.0
    std_dev_raw = math.sqrt(variance_raw)

    if n >= 4:
        try:
            quantiles = statistics.quantiles(sorted_scores, n=4)
            q1_raw, q3_raw = quantiles[0], quantiles[2]
        except statistics.StatisticsError:
            q1_raw, q3_raw = min_raw, max_raw
    else:
        # Quartiles are not meaningfully defined for very small samples, so the
        # observed range is reported rather than a misleading interpolated value.
        q1_raw, q3_raw = min_raw, max_raw

    return {
        "count": n,
        "mean": round(statistics.fmean(clean_scores), 2),
        "median": round(statistics.median(clean_scores), 2),
        "std_dev": round(std_dev_raw, 2),
        "variance": round(variance_raw, 2),
        "min_score": round(min_raw, 2),
        "max_score": round(max_raw, 2),
        "q1": round(q1_raw, 2),
        "q3": round(q3_raw, 2),
        # The true interquartile range, not the difference of the two rounded
        # quartiles above.
        "iqr": round(q3_raw - q1_raw, 2),
        "qualification_rate": ats_config.qualification_rate_for(clean_scores),
        "distribution_buckets": ats_config.stratify_scores(clean_scores),
    }


def percentage(numerator, denominator, cap_at_100=True):
    """
    A ratio as a percentage, safe against a zero or missing denominator.

    ``cap_at_100`` clamps the result to 100 because several of the funnel rates
    compare two independently counted populations (for example letters issued
    against clearances granted); a ratio above 100 means the populations overlap
    rather than that the rate exceeded a ceiling, and reporting 112% would read
    as an impossible figure.
    """
    if not denominator:
        return 0.0
    value = (numerator / denominator) * 100
    if cap_at_100:
        value = min(value, 100.0)
    return round(value, 1)


def calculate_conversion_funnel(apps_qs, deployments_qs, clearances_qs, letters_count):
    """
    Computes end-to-end actuarial yields, throughput, and attrition rates.

    The funnel follows the binary selection model. There is no shortlisting
    stage: an application is either decided SUCCESSFUL or REJECTED, so the
    previous "Shortlisted -> Offers & Hired" pair of stages collapsed into a
    single "Successful Selections" stage and a "Rejected" stage that accounts
    for the decisions taken against each intake.

    Two guarantees hold for every rate:

    * ``rate`` is always the stage's share of *all applications*, so the column
      can be read straight down as a funnel. It previously mixed denominators
      (stages 2-3 used all applications, stages 4-6 used the previous stage), so
      the numbers in one column were not comparable.
    * ``stage_conversion`` is the stage-over-stage rate, which is what the
      ``benchmark`` text describes.
    """
    total_apps = apps_qs.count()
    successful_count = apps_qs.filter(status="SUCCESSFUL").count()
    rejected_count = apps_qs.filter(status="REJECTED").count()
    # PENDING and REVIEWED are workflow markers, not outcomes, so "awaiting
    # decision" is what they mean here.
    pending_count = apps_qs.filter(status__in=Application.PENDING_STATUSES).count()

    total_deployments = deployments_qs.count()
    active_deployments = deployments_qs.filter(status="DEPLOYED").count()
    exited_deployments = deployments_qs.filter(status="EXITED").count()

    total_clearances = clearances_qs.count()
    cleared_count = clearances_qs.filter(status="CLEARED").count()
    pending_clearances = clearances_qs.filter(status__in=["PENDING_DEPARTMENT", "PENDING_HR"]).count()
    rejected_clearances = clearances_qs.filter(status="REJECTED").count()

    # Share of all applicants who were granted the opportunity.
    success_rate = percentage(successful_count, total_apps)
    rejection_rate = percentage(rejected_count, total_apps)
    # Of the applications actually decided, how many were successful. The
    # denominator is the decided total, not all applications, so an intake with
    # many undecided applications cannot depress this figure.
    decided_count = successful_count + rejected_count
    success_among_decided_rate = percentage(successful_count, decided_count)
    deployment_yield_rate = percentage(active_deployments, successful_count)
    clearance_completion_rate = percentage(cleared_count, total_clearances)
    letter_issuance_rate = percentage(letters_count, cleared_count)

    def stage(count, previous_count, benchmark_label):
        return {
            "stage_conversion": percentage(count, previous_count),
            "rate": percentage(count, total_apps),
            "benchmark": benchmark_label,
        }

    return {
        "funnel_stages": [
            {
                "stage": "1. Applications Submitted",
                "count": total_apps,
                "rate": 100.0 if total_apps else 0.0,
                "stage_conversion": 100.0 if total_apps else 0.0,
                "benchmark": "Baseline 100%",
            },
            {
                "stage": "2. Successful Selections",
                "count": successful_count,
                **stage(successful_count, total_apps, f"{success_rate}% of applicants"),
            },
            {
                "stage": "3. Rejected",
                "count": rejected_count,
                **stage(rejected_count, total_apps, f"{rejection_rate}% of applicants"),
            },
            {
                "stage": "4. Active Deployments",
                "count": active_deployments,
                **stage(active_deployments, successful_count, f"{deployment_yield_rate}% of successful"),
            },
            {
                "stage": "5. Dual Cleared",
                "count": cleared_count,
                **stage(cleared_count, total_clearances, f"{clearance_completion_rate}% of clearances"),
            },
            {
                "stage": "6. Letters Issued",
                "count": letters_count,
                **stage(letters_count, cleared_count, f"{letter_issuance_rate}% of cleared"),
            },
        ],
        "metrics": {
            "success_rate": success_rate,
            "rejection_rate": rejection_rate,
            "success_among_decided_rate": success_among_decided_rate,
            # Retained name: the Admin analytics card labels this "Placement
            # Yield", and the value is unchanged in meaning (share of all
            # applicants granted the opportunity).
            "placement_yield_rate": success_rate,
            "deployment_yield_rate": deployment_yield_rate,
            "clearance_completion_rate": clearance_completion_rate,
            "letter_issuance_rate": letter_issuance_rate,
            "total_applications": total_apps,
            "successful_applications": successful_count,
            "decided_applications": decided_count,
            "total_deployments": total_deployments,
            "exited_deployments": exited_deployments,
            "total_clearances": total_clearances,
            "pending_review_count": pending_count,
            "rejected_count": rejected_count,
            "pending_clearances_count": pending_clearances,
            "rejected_clearances_count": rejected_clearances,
        },
    }


class AnalyticsView(APIView):
    """Programme analytics.

    Admin/HR receive the system-wide institutional payload.  Directors receive a
    payload built exclusively from their own department's records.  Applicants
    are not permitted: the system-wide payload contains cross-department
    capacity figures and audit-trail descriptions.
    """

    def get_permissions(self):
        return [IsDirectorOrHRorAdmin()]

    def get(self, request):
        now = timezone.now()
        user = request.user
        start_date_str = request.query_params.get("start_date")
        end_date_str = request.query_params.get("end_date")

        # A single supplied bound used to be discarded entirely, silently
        # replacing the caller's request with a default 30-day window -- the
        # filter looked applied but was not. Each bound is now honoured on its
        # own, and an unparseable value is rejected rather than ignored.
        default_window = timedelta(days=30)
        window_error = None
        parsed_start = parsed_end = None
        for raw, label in ((start_date_str, "start_date"), (end_date_str, "end_date")):
            if not raw:
                continue
            try:
                parsed = datetime.strptime(raw, "%Y-%m-%d")
            except ValueError:
                window_error = f"'{raw}' is not a valid {label}. Use the YYYY-MM-DD format."
                break
            if label == "start_date":
                parsed_start = parsed
            else:
                # The whole end day is included. The old code stopped at
                # 23:59:59.000000, which excluded the final second of the day.
                parsed_end = parsed + timedelta(days=1) - timedelta(microseconds=1)

        if window_error:
            return Response({"detail": window_error}, status=status.HTTP_400_BAD_REQUEST)

        start_date = (
            timezone.make_aware(parsed_start)
            if parsed_start is not None
            else now - default_window
        )
        end_date = timezone.make_aware(parsed_end) if parsed_end is not None else now

        if start_date > end_date:
            return Response(
                {"detail": "'start_date' cannot be later than 'end_date'."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        delta_days = (end_date - start_date).days

        if delta_days <= 45:
            trunc_func = TruncDate("applied_at")
            date_format = "%b %d"
            period_label = "Daily"
        elif delta_days <= 180:
            trunc_func = TruncWeek("applied_at")
            date_format = "%b %d"
            period_label = "Weekly"
        elif delta_days <= 730:
            trunc_func = TruncMonth("applied_at")
            date_format = "%b %Y"
            period_label = "Monthly"
        else:
            trunc_func = TruncYear("applied_at")
            date_format = "%Y"
            period_label = "Yearly"

        # Check for Department Director scope
        if is_director_role(user):
            dept = get_director_department(user)

            dept_jobs = Job.objects.filter(department=dept, is_archived=False) if dept else Job.objects.none()
            dept_apps = Application.objects.filter(job__department=dept) if dept else Application.objects.none()
            dept_reqs = Requisition.objects.filter(department=dept) if dept else Requisition.objects.none()
            dept_deployments = Deployment.objects.filter(department=dept) if dept else Deployment.objects.none()
            dept_clearances = Clearance.objects.filter(deployment__department=dept) if dept else Clearance.objects.none()

            total_slots_req, total_slots_fil = job_capacity_totals(dept_jobs)
            available_headroom = max(total_slots_req - total_slots_fil, 0)
            dept_fill_rate = percentage(total_slots_fil, total_slots_req)

            active_attachees_qs = dept_deployments.filter(status="DEPLOYED").select_related("application__user")[:10]
            active_attachees_list = [
                {
                    "id": d.id,
                    "name": attachee_display_name(d.application.user),
                    "email": d.application.user.email,
                    "start_date": str(d.start_date),
                    "planned_end_date": str(d.planned_end_date),
                    "status": d.status,
                }
                for d in active_attachees_qs
            ]

            pending_clearance_qs = dept_clearances.filter(status="PENDING_DEPARTMENT").select_related("deployment__application__user")[:10]
            pending_clearances_list = [
                {
                    "id": c.id,
                    # Written as an explicit helper because the original inline
                    # expression relied on `or` binding looser than the
                    # conditional, which is correct but very easy to break.
                    "name": clearance_attachee_name(c),
                    "submitted_at": c.final_report_submitted_at.isoformat() if c.final_report_submitted_at else None,
                    "has_report": bool(c.final_report_file),
                }
                for c in pending_clearance_qs
            ]

            # Department statistical quality metrics
            dept_scores = list(dept_apps.exclude(ats_score__isnull=True).values_list("ats_score", flat=True))
            dept_stats = calculate_statistical_metrics(dept_scores)

            pending_stage1_count = dept_clearances.filter(status="PENDING_DEPARTMENT").count()
            # Velocity is the share of stage-1 decisions that were approvals, so
            # the denominator must be the clearances that actually reached
            # stage 1. It previously counted every clearance in existence while
            # the numerator counted only approved ones, so a department with
            # many untouched NOT_STARTED clearances had its velocity understated
            # for no reason.
            stage1_decided = dept_clearances.exclude(status="NOT_STARTED")
            approved_stage1_count = stage1_decided.filter(department_cleared=True).count()
            stage1_decided_count = stage1_decided.count()
            stage1_velocity_rate = percentage(approved_stage1_count, stage1_decided_count)

            return Response(
                {
                    "role": "DEPARTMENT_DIRECTOR",
                    "department_info": {
                        "id": dept.id,
                        "name": dept.name,
                        "description": dept.description,
                        "location": dept.location,
                        "typical_intake_capacity": dept.typical_intake_capacity or 0,
                    } if dept else None,
                    "total_vacancies": dept_jobs.count(),
                    "total_slots_required": total_slots_req,
                    "total_slots_filled": total_slots_fil,
                    "available_headroom": available_headroom,
                    "fill_rate": dept_fill_rate,
                    "total_applications": dept_apps.count(),
                    "average_ats_score": dept_stats["mean"],
                    "statistical_quality_metrics": dept_stats,
                    # Sent so the dashboard labels the benchmark it actually
                    # measured against instead of hardcoding a figure that
                    # stops matching as soon as an Admin re-configures it.
                    "qualification_benchmark": float(
                        AtsScoringConfiguration.get_config().qualification_benchmark
                    ),
                    "active_attachees_count": dept_deployments.filter(status="DEPLOYED").count(),
                    "exited_attachees_count": dept_deployments.filter(status="EXITED").count(),
                    "pending_stage1_clearances": pending_stage1_count,
                    "approved_stage1_clearances": approved_stage1_count,
                    "decided_stage1_clearances": stage1_decided_count,
                    "stage1_velocity_rate": stage1_velocity_rate,
                    "requisitions_summary": {
                        "total": dept_reqs.count(),
                        "pending": dept_reqs.filter(status="PENDING").count(),
                        "approved": dept_reqs.filter(status="APPROVED").count(),
                        "fulfilled": dept_reqs.filter(status="FULFILLED").count(),
                        "rejected": dept_reqs.filter(status="REJECTED").count(),
                        "candidates_requested": dept_reqs.aggregate(
                            total=Sum("num_candidates_requested")
                        )["total"] or 0,
                    },
                    "active_attachees": active_attachees_list,
                    "pending_clearances": pending_clearances_list,
                }
            )

        # Admin & HR System-Wide Analytics
        apps = Application.objects.filter(applied_at__gte=start_date, applied_at__lte=end_date)
        all_apps = Application.objects.all()
        jobs_qs = Job.objects.filter(is_archived=False)

        total_jobs = jobs_qs.count()
        total_applications = apps.count()

        status_distribution = list(apps.values("status").annotate(count=Count("id")))
        job_type_distribution = list(
            apps.values("job__job_type").annotate(count=Count("id"))
        )

        formatted_job_type_dict = []
        for j in job_type_distribution:
            formatted_job_type_dict.append(
                {"job_type": j["job__job_type"], "count": j["count"]}
            )

        # Statistical Quality Metrics Engine
        period_scores = list(apps.exclude(ats_score__isnull=True).values_list("ats_score", flat=True))
        stats_metrics = calculate_statistical_metrics(period_scores)

        trend_qs = (
            apps.annotate(period_date=trunc_func)
            .values("period_date")
            .annotate(app_count=Count("id"), avg_score=Avg("ats_score"))
            .order_by("period_date")
        )

        trend = []
        for item in trend_qs:
            trend.append(
                {
                    "date": (
                        item["period_date"].strftime(date_format)
                        if hasattr(item["period_date"], "strftime")
                        else str(item["period_date"])
                    ),
                    "applications": item["app_count"],
                    "average_score": (
                        round(item["avg_score"], 2) if item["avg_score"] else 0.0
                    ),
                }
            )

        # Operational Aggregates
        total_users = User.objects.count()
        users_summary = {
            "total_users": total_users,
            "applicants": User.objects.filter(role="APPLICANT").count(),
            "hr": User.objects.filter(role="HR").count(),
            "directors": User.objects.filter(role__in=["DEPARTMENT_DIRECTOR", "DEPARTMENT"]).count(),
            "admins": User.objects.filter(role="ADMIN").count(),
        }

        total_capacity, total_filled = job_capacity_totals(jobs_qs)
        overall_fill_rate = percentage(total_filled, total_capacity)
        available_headroom = max(total_capacity - total_filled, 0)

        departments_summary = {
            "total": Department.objects.filter(is_archived=False).count(),
            "active": Department.objects.filter(is_active=True, is_archived=False).count(),
            "total_capacity": total_capacity,
            "total_filled": total_filled,
            "available_headroom": available_headroom,
            "overall_fill_rate": overall_fill_rate,
        }

        total_reqs = Requisition.objects.count()
        approved_reqs = Requisition.objects.filter(status__in=["APPROVED", "FULFILLED"]).count()
        req_approval_rate = percentage(approved_reqs, total_reqs)

        reqs_approval_numerator = Requisition.objects.filter(status__in=["APPROVED", "FULFILLED"]).count()
        requisitions_summary = {
            "total": total_reqs,
            "pending": Requisition.objects.filter(status="PENDING").count(),
            "approved": Requisition.objects.filter(status="APPROVED").count(),
            "fulfilled": Requisition.objects.filter(status="FULFILLED").count(),
            "rejected": Requisition.objects.filter(status="REJECTED").count(),
            "candidates_requested": Requisition.objects.aggregate(
                total=Sum("num_candidates_requested")
            )["total"] or 0,
            "approval_rate": percentage(reqs_approval_numerator, total_reqs),
        }

        deployments_qs = Deployment.objects.all()
        deployments_summary = {
            "total": deployments_qs.count(),
            "active": deployments_qs.filter(status="DEPLOYED").count(),
            "exited": deployments_qs.filter(status="EXITED").count(),
            "pending_deployment": deployments_qs.filter(status="PENDING_DEPLOYMENT").count(),
        }

        clearances_qs = Clearance.objects.all()
        clearances_summary = {
            "total": clearances_qs.count(),
            "pending_department": clearances_qs.filter(status="PENDING_DEPARTMENT").count(),
            "pending_hr": clearances_qs.filter(status="PENDING_HR").count(),
            "cleared": clearances_qs.filter(status="CLEARED").count(),
            "rejected": clearances_qs.filter(status="REJECTED").count(),
        }

        letters_count = RecommendationLetter.objects.count()
        recommendations_summary = {
            "total_issued": letters_count,
        }

        # End-to-End Actuarial Conversion Funnel
        conversion_funnel = calculate_conversion_funnel(all_apps, deployments_qs, clearances_qs, letters_count)

        # Department Fill Rates breakdown
        dept_fill_rates = []
        for d in Department.objects.filter(is_archived=False, is_active=True).prefetch_related("jobs", "deployments"):
            d_jobs = d.jobs.filter(is_archived=False)
            # No fallback to typical_intake_capacity here. A department with no
            # open vacancies has no capacity to fill, and substituting its
            # intake figure reported phantom slots, a 0% fill rate and headroom
            # for a department that had posted nothing. It also made the
            # per-department slots_required values fail to sum to the
            # system-wide total_capacity computed above.
            d_req, d_fil = job_capacity_totals(d_jobs)
            d_active_attachees = d.deployments.filter(status="DEPLOYED").count()
            d_headroom = max(d_req - d_fil, 0)
            dept_fill_rates.append({
                "id": d.id,
                "name": d.name,
                "slots_required": d_req,
                "slots_filled": d_fil,
                "available_headroom": d_headroom,
                "fill_rate": percentage(d_fil, d_req),
                "active_attachees": d_active_attachees,
            })

        # Recent Activity from Audit Logs (using exact AuditLog model fields: actor & target_description)
        recent_activity = []
        try:
            from accounts.models import AuditLog
            logs = AuditLog.objects.select_related("actor").order_by("-created_at")[:10]
            for log in logs:
                recent_activity.append({
                    "id": log.id,
                    "action": log.action,
                    "description": log.target_description,
                    "created_at": log.created_at.isoformat(),
                    "user": log.actor.username if log.actor else "System",
                })
        except Exception:
            pass

        return Response(
            {
                "role": user.role,
                "total_jobs": total_jobs,
                "total_applications": total_applications,
                "status_distribution": status_distribution,
                "job_type_distribution": formatted_job_type_dict,
                "average_ats_score": stats_metrics["mean"],
                "statistical_quality_metrics": stats_metrics,
                # The benchmark the qualification rate was actually computed
                # against, so the dashboard does not label a stale hardcoded
                # figure after an Admin changes the configuration.
                "qualification_benchmark": float(
                    AtsScoringConfiguration.get_config().qualification_benchmark
                ),
                "conversion_funnel": conversion_funnel,
                "trend": trend,
                "period_label": period_label,
                "users_summary": users_summary,
                "departments_summary": departments_summary,
                "requisitions_summary": requisitions_summary,
                "deployments_summary": deployments_summary,
                "clearances_summary": clearances_summary,
                "recommendations_summary": recommendations_summary,
                "department_fill_rates": dept_fill_rates,
                "recent_activity": recent_activity,
            }
        )


class DepartmentStaffingView(APIView):
    """
    Per-department staffing figures, computed once on the server.

    The Staffing Overview used to be assembled in the browser from the
    departments, requisitions and deployments list endpoints, which produced two
    figures that were wrong rather than merely stale: it counted a Deployment
    status (``ACTIVE``) that does not exist, so the active column was always
    zero, and it divided by the hand-edited ``typical_intake_capacity`` rather
    than by open advertised capacity, so departments with no vacancies at all
    still reported slots and headroom. See ``jobs.staffing`` for the full
    account.

    Serving it from here also means the Overview, the fill-rate report and the
    analytics dashboards are all reading the same aggregation rather than three
    implementations of it.

    Scoped exactly as AnalyticsView: Admin/HR see every department, a Director
    sees only the department they are bound to server-side, and applicants are
    refused because the payload is cross-department capacity data.
    """

    def get_permissions(self):
        return [IsDirectorOrHRorAdmin()]

    def get(self, request):
        from .staffing import clearance_bucket_counts, department_staffing_rows, staffing_summary

        department_ids = None
        if not (request.user.role in ["ADMIN", "HR"]):
            department = get_director_department(request.user)
            # A Director with no department has no staffing to see. An empty
            # list is returned rather than the unscoped system-wide payload.
            department_ids = [department.id] if department else []

        rows = department_staffing_rows(department_ids)
        # Reuse the rows already fetched rather than making the summary helper
        # re-issue the same nine aggregate queries.
        summary = staffing_summary(rows=rows)
        return Response(
            {
                "role": request.user.role,
                "scoped_to_department": department_ids is not None
                and len(department_ids) == 1,
                "summary": summary,
                "departments": rows,
                "clearance": clearance_bucket_counts(department_ids),
            }
        )


class EligibilitySettingView(APIView):
    """
    Public, read-only GET endpoint returning the general institutional attachment
    eligibility requirements as sanitized rich text.

    Admin-only PATCH endpoint. This is the sole write path for the content; the
    public landing page is a pure read-only display and exposes no edit control.
    """
    CACHE_TIMEOUT = 900  # 15 minutes, matching the public vacancy listing cache

    def get_permissions(self):
        if self.request.method in permissions.SAFE_METHODS:
            return [permissions.AllowAny()]
        return [IsAdminOnly()]

    def get(self, request):
        from django.core.cache import cache
        from .signals import CACHE_KEY_ELIGIBILITY_SETTINGS

        # Read-through cache. The cache key is invalidated by the
        # EligibilitySetting post_save/post_delete signal in jobs/signals.py, so
        # an Admin edit is visible on the landing page's very next load.
        data = cache.get(CACHE_KEY_ELIGIBILITY_SETTINGS)
        if data is None:
            settings_obj = EligibilitySetting.get_settings()
            data = EligibilitySettingSerializer(settings_obj).data
            cache.set(CACHE_KEY_ELIGIBILITY_SETTINGS, data, self.CACHE_TIMEOUT)
        return Response(data, status=status.HTTP_200_OK)

    def patch(self, request):
        from django.core.cache import cache
        from accounts.audit import log_audit_event
        from .signals import CACHE_KEY_ELIGIBILITY_SETTINGS

        settings_obj = EligibilitySetting.get_settings()
        serializer = EligibilitySettingSerializer(settings_obj, data=request.data, partial=True)
        serializer.is_valid(raise_exception=True)
        serializer.save(updated_by=request.user)
        log_audit_event(
            request,
            "ELIGIBILITY_CONTENT_UPDATED",
            "Administrator updated the public institutional eligibility requirements content",
            metadata={"setting_id": settings_obj.id},
        )
        # Belt and braces: the post_save signal also clears this key, but an
        # explicit delete keeps the public response correct even if the signal
        # is ever unregistered.
        cache.delete(CACHE_KEY_ELIGIBILITY_SETTINGS)
        return Response(serializer.data, status=status.HTTP_200_OK)


class AtsScoringConfigurationView(APIView):
    """
    Admin-only read/write access to the ATS quality tier boundaries.

    Every tier shown on the Admin dashboard's "Applicant Quality Distribution &
    Statistical Dispersion Profile" card is derived from these values via
    AtsScoringConfiguration.tier_definitions() / stratify_scores(); changing them
    re-tiers existing applicants on the next dashboard load with no re-scoring
    and no redeploy.
    """
    permission_classes = [IsAdminOnly]

    def get(self, request):
        config = AtsScoringConfiguration.get_config()
        return Response(AtsScoringConfigurationSerializer(config).data, status=status.HTTP_200_OK)

    def patch(self, request):
        from accounts.audit import log_audit_event

        config = AtsScoringConfiguration.get_config()
        serializer = AtsScoringConfigurationSerializer(config, data=request.data, partial=True)
        serializer.is_valid(raise_exception=True)
        serializer.save(updated_by=request.user)
        log_audit_event(
            request,
            "ATS_SCORING_CONFIG_UPDATED",
            "Administrator updated the ATS quality tier scoring thresholds",
            metadata={
                "exceptional_min": str(serializer.validated_data.get("exceptional_min", config.exceptional_min)),
                "highly_qualified_min": str(serializer.validated_data.get("highly_qualified_min", config.highly_qualified_min)),
                "qualified_min": str(serializer.validated_data.get("qualified_min", config.qualified_min)),
                "average_min": str(serializer.validated_data.get("average_min", config.average_min)),
                "qualification_benchmark": str(serializer.validated_data.get("qualification_benchmark", config.qualification_benchmark)),
            },
        )
        return Response(serializer.data, status=status.HTTP_200_OK)


class ArchivesListView(APIView):
    """
    Consolidated archive explorer.

    Admin/HR: the full system archive — archived vacancies, applications and
    department profiles, with optional type / department / search filtering.
    These users are also the only ones who can restore records (see
    ``ArchiveRestoreView``) or purge them (see ``ArchivePurgeView``).

    Directors: a strictly read-only view of their OWN department's archived
    vacancies and applications.  The department scope is derived from the
    authenticated user and is never read from client input; asking for another
    department is rejected with 403, and asking for archived department
    profiles (an institutional record) is rejected with 403.  The returned
    counts are scoped the same way as the records, so the summary tiles cannot
    leak system-wide archive totals.
    """

    def get_permissions(self):
        return [IsDirectorOrHRorAdmin()]

    def get(self, request):
        entity_type = request.query_params.get("type", "ALL").upper()  # 'ALL', 'VACANCY', 'APPLICATION', 'DEPARTMENT'
        dept_id = request.query_params.get("department")
        search = request.query_params.get("search", "").strip()
        date_from_raw = request.query_params.get("date_from", "").strip()
        date_to_raw = request.query_params.get("date_to", "").strip()

        # ── Date range validation ───────────────────────────────────────────
        # A silently ignored filter is worse than a rejected one: an Admin who
        # mistypes a date would otherwise believe they had filtered the archive.
        date_from = self._parse_archive_date(date_from_raw, "date_from")
        date_to = self._parse_archive_date(date_to_raw, "date_to")
        if date_from and date_to and date_from > date_to:
            return Response(
                {"detail": "'Date From' cannot be later than 'Date To'."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        def archive_window(qs, *fallback_chain):
            """
            Restricts a queryset to the effective archive date, i.e. the same
            value the record reports as ``archived_at`` (which falls back to
            creation/update stamps when no explicit archive date was recorded).

            The Coalesce matters: a legacy archived row with a NULL
            ``archived_at`` would otherwise be excluded by every date filter
            while still displaying a creation date that is plainly in range,
            which reads as "the filter is broken".

            Accepts a chain of fields, most specific first, so a completed
            attachee can be dated by its actual HR clearance decision rather
            than by whenever the archive flag happened to be written.

            The date is projected into a concrete annotation first because a
            Coalesce expression cannot be used as the left-hand side of a
            ``__date__gte`` lookup through a ``**kwargs`` lookup key.
            """
            if not date_from and not date_to:
                return qs
            if len(fallback_chain) == 1:
                # The queryset already exposes a resolved completion date, so
                # re-wrapping it in a Coalesce would just restate its own
                # internal fallbacks.
                qs = qs.annotate(_effective_archive_date=TruncDate(fallback_chain[0]))
            else:
                qs = qs.annotate(
                    _effective_archive_date=TruncDate(Coalesce(*fallback_chain))
                )
            if date_from:
                qs = qs.filter(_effective_archive_date__gte=date_from)
            if date_to:
                qs = qs.filter(_effective_archive_date__lte=date_to)
            return qs

        director_view = is_director_role(request.user)
        director_department = get_director_department(request.user) if director_view else None

        if director_view:
            if director_department is None:
                # A Director with no department assigned sees an empty archive,
                # never an unscoped one.
                return Response(
                    {
                        "counts": {"total": 0, "vacancies": 0, "applications": 0, "attachments": 0, "departments": 0},
                        "results": [],
                        "department": None,
                        "read_only": True,
                    },
                    status=status.HTTP_200_OK,
                )
            if dept_id and str(director_department.id) != str(dept_id):
                return Response(
                    {"detail": "You can only view archived records for your own assigned department."},
                    status=status.HTTP_403_FORBIDDEN,
                )
            if entity_type == "DEPARTMENT":
                return Response(
                    {"detail": "Archived department profiles are an institutional record and are not available to department directors."},
                    status=status.HTTP_403_FORBIDDEN,
                )
            # Locked to the Director's own department for the rest of the request.
            dept_id = str(director_department.id)
        elif dept_id == "ALL":
            dept_id = None

        archived_items = []

        # 1. Archived Vacancies
        job_qs = None
        if entity_type in ["ALL", "VACANCY"]:
            job_qs = Job.objects.filter(is_archived=True).select_related("department", "archived_by")
            if dept_id and dept_id != "ALL":
                job_qs = job_qs.filter(department_id=dept_id)
            if search:
                job_qs = job_qs.filter(Q(title__icontains=search) | Q(department__name__icontains=search) | Q(description__icontains=search))
            job_qs = archive_window(job_qs, "archived_at", "created_at")
            for j in job_qs:
                archived_items.append({
                    "id": j.id,
                    "entity_type": "VACANCY",
                    "title": j.title,
                    "department_id": j.department_id,
                    "department_name": j.department.name if j.department else "General",
                    "slots": j.slots_required,
                    "archived_at": j.archived_at.isoformat() if j.archived_at else j.created_at.isoformat(),
                    "archived_by": j.archived_by.username if j.archived_by else "System",
                    "details": f"Attachment Vacancy ({j.slots_required} slots, {j.duration_weeks} weeks)",
                })

        # 2. Archived Applications
        #
        # A successfully completed attachment is reported as its own ATTACHMENT
        # record instead, so it is excluded here. Without this exclusion the
        # same attachee would be listed twice — once as a plain application and
        # once as a completed attachee — and the counts would double-count them.
        app_qs = None
        if entity_type in ["ALL", "APPLICATION"]:
            app_qs = Application.objects.filter(is_archived=True).exclude(
                COMPLETED_ATTACHMENT_CONDITION
            ).select_related("job", "job__department", "user", "archived_by")
            if dept_id and dept_id != "ALL":
                app_qs = app_qs.filter(job__department_id=dept_id)
            if search:
                app_qs = app_qs.filter(Q(user__username__icontains=search) | Q(user__first_name__icontains=search) | Q(user__last_name__icontains=search) | Q(job__title__icontains=search))
            app_qs = archive_window(app_qs, "archived_at", "applied_at")
            for a in app_qs:
                archived_items.append({
                    "id": a.id,
                    "entity_type": "APPLICATION",
                    "title": f"{a.user.get_full_name() or a.user.username} - {a.job.title if a.job else 'General Application'}",
                    "department_id": a.job.department_id if a.job else None,
                    "department_name": a.job.department.name if a.job and a.job.department else "General",
                    "slots": None,
                    "archived_at": a.archived_at.isoformat() if a.archived_at else a.applied_at.isoformat(),
                    "archived_by": a.archived_by.username if a.archived_by else "System",
                    "details": f"Status: {a.get_status_display()}, ATS Score: {a.ats_score}%",
                })

        # 3. Completed attachments — the attachees who finished here long ago and
        #    are kept as a referable record. Available to Directors, but scoped
        #    to their own department exactly like every other Director read.
        att_qs = None
        if entity_type in ["ALL", "ATTACHMENT"]:
            att_qs = completed_attachment_queryset()
            if dept_id and dept_id != "ALL":
                att_qs = att_qs.filter(job__department_id=dept_id)
            if search:
                att_qs = att_qs.filter(
                    Q(user__username__icontains=search)
                    | Q(user__first_name__icontains=search)
                    | Q(user__last_name__icontains=search)
                    | Q(user__email__icontains=search)
                    | Q(user__profile__id_number__icontains=search)
                    | Q(user__profile__institution_name__icontains=search)
                    | Q(job__title__icontains=search)
                    | Q(job__department__name__icontains=search)
                )
            # ``_completion_date`` is already the coalesced "when did this
            # attachment end" value, so it is passed as the sole chain entry.
            att_qs = archive_window(att_qs, "_completion_date")
            for a in att_qs:
                archived_items.append(serialize_attachment_archive_row(a))

        # 4. Archived Departments — institutional records, never exposed to Directors.
        dept_qs = None
        if entity_type in ["ALL", "DEPARTMENT"] and not director_view:
            dept_qs = Department.objects.filter(is_archived=True).select_related("director", "archived_by")
            if search:
                dept_qs = dept_qs.filter(Q(name__icontains=search) | Q(location__icontains=search))
            dept_qs = archive_window(dept_qs, "archived_at", "updated_at")
            for d in dept_qs:
                archived_items.append({
                    "id": d.id,
                    "entity_type": "DEPARTMENT",
                    "title": d.name,
                    "department_id": d.id,
                    "department_name": d.name,
                    "slots": d.typical_intake_capacity,
                    "archived_at": d.archived_at.isoformat() if d.archived_at else d.updated_at.isoformat(),
                    "archived_by": d.archived_by.username if d.archived_by else "System",
                    "details": f"Location: {d.location or 'Nairobi'}, Director: {d.director.username if d.director else 'Unassigned'}",
                })

        archived_items.sort(key=lambda x: x["archived_at"] or "", reverse=True)

        # Counts are derived from the SAME filtered querysets that produced the
        # records above, so the summary tiles always agree with the table below
        # them. Previously `total` was filtered while the three breakdown counts
        # were unconditional, which meant an Admin could see "3 records found"
        # next to "120 archived vacancies" and have no way to tell which was
        # telling the truth.
        #
        # Director safety is unchanged: a Director's querysets were already
        # locked to their own department above, and archived department profiles
        # remain excluded entirely, so no count can ever expose another
        # department's archive activity.
        counts = {
            "vacancies": job_qs.count() if job_qs is not None else 0,
            "applications": app_qs.count() if app_qs is not None else 0,
            "attachments": att_qs.count() if att_qs is not None else 0,
            "departments": dept_qs.count() if dept_qs is not None else 0,
        }
        counts["total"] = (
            counts["vacancies"]
            + counts["applications"]
            + counts["attachments"]
            + counts["departments"]
        )

        department_payload = None
        if director_view:
            department_payload = {
                "id": director_department.id,
                "name": director_department.name,
            }

        return Response(
            {
                "counts": counts,
                "results": archived_items,
                "department": department_payload,
                "read_only": director_view,
                "applied_filters": {
                    "type": entity_type,
                    "department": dept_id if dept_id and dept_id != "ALL" else None,
                    "search": search or None,
                    "date_from": date_from.isoformat() if date_from else None,
                    "date_to": date_to.isoformat() if date_to else None,
                },
            },
            status=status.HTTP_200_OK,
        )

    @staticmethod
    def _parse_archive_date(raw, field_name):
        """
        Parses a ``YYYY-MM-DD`` filter value into a ``date``.

        Returns ``None`` for an empty value, and raises ``ValidationError`` for
        anything unparseable so a typo surfaces as a clear 400 instead of being
        silently dropped.
        """
        if not raw:
            return None
        try:
            return datetime.strptime(raw, "%Y-%m-%d").date()
        except ValueError:
            raise ValidationError({field_name: f"'{raw}' is not a valid date. Use the YYYY-MM-DD format."})


class ArchiveRestoreView(APIView):
    """
    Restores an archived entity to active status. Accessible to Admin & HR.
    """
    permission_classes = [IsHRorAdmin]

    def post(self, request, entity_type, pk):
        from accounts.audit import log_audit_event
        entity_type = entity_type.upper()

        if entity_type in ["VACANCY", "JOB"]:
            try:
                item = Job.objects.get(pk=pk)
                item.is_archived = False
                # Reactivation is part of the restore, not an optional extra.
                # Every archive path withdraws the vacancy by clearing
                # `is_active` as well, and the public and applicant listings both
                # filter on `is_active=True`. Clearing only `is_archived` left a
                # half-restored vacancy that HR and Admin saw in the Active tab
                # while no applicant could ever see it, because they were served a
                # list filtered on `is_active=True`.
                item.is_active = True
                item.archived_at = None
                item.archived_by = None
                item.save(update_fields=["is_archived", "is_active", "archived_at", "archived_by"])
                log_audit_event(
                    request,
                    "RECORD_RESTORED",
                    f"Restored archived vacancy '{item.title}' (ID: {item.id}) to active listings",
                    metadata={"entity_type": "VACANCY", "id": item.id, "title": item.title}
                )
                return Response({"detail": f"Vacancy '{item.title}' has been restored to active status."})
            except Job.DoesNotExist:
                return Response({"detail": "Archived vacancy not found."}, status=status.HTTP_404_NOT_FOUND)

        elif entity_type == "APPLICATION":
            try:
                item = Application.objects.get(pk=pk)
                item.is_archived = False
                item.archived_at = None
                item.archived_by = None
                item.save(update_fields=["is_archived", "archived_at", "archived_by"])
                log_audit_event(
                    request,
                    "RECORD_RESTORED",
                    f"Restored archived application #{item.id} ({item.user.username}) to active applications",
                    metadata={"entity_type": "APPLICATION", "id": item.id, "applicant": item.user.username}
                )
                return Response({"detail": f"Application #{item.id} has been restored to active status."})
            except Application.DoesNotExist:
                return Response({"detail": "Archived application not found."}, status=status.HTTP_404_NOT_FOUND)

        elif entity_type == "ATTACHMENT":
            # The archive stores completed attachees as their underlying
            # application, so a restore just clears the archive flags again.
            # The attachee stays in the ATTACHMENT view regardless, because
            # completion — not the flag — is what makes them an archival record.
            try:
                item = Application.objects.select_related("user", "job").get(pk=pk)
            except Application.DoesNotExist:
                return Response({"detail": "Archived attachee record not found."}, status=status.HTTP_404_NOT_FOUND)
            if not is_completed_attachment(item):
                return Response(
                    {"detail": "This application is not a completed attachment, so it cannot be restored as an attachee record."},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            item.is_archived = False
            item.archived_at = None
            item.archived_by = None
            item.save(update_fields=["is_archived", "archived_at", "archived_by"])
            log_audit_event(
                request,
                "RECORD_RESTORED",
                f"Restored archived attachee record #{item.id} ({item.user.username})",
                metadata={"entity_type": "ATTACHMENT", "id": item.id, "applicant": item.user.username}
            )
            return Response({"detail": f"Attachee record for {item.user.username} has been restored."})

        elif entity_type == "DEPARTMENT":
            try:
                item = Department.objects.get(pk=pk)
                item.is_archived = False
                item.archived_at = None
                item.archived_by = None
                item.save(update_fields=["is_archived", "archived_at", "archived_by"])
                log_audit_event(
                    request,
                    "RECORD_RESTORED",
                    f"Restored archived department '{item.name}' (ID: {item.id}) to active directory",
                    metadata={"entity_type": "DEPARTMENT", "id": item.id, "name": item.name}
                )
                return Response({"detail": f"Department '{item.name}' has been restored to active status."})
            except Department.DoesNotExist:
                return Response({"detail": "Archived department not found."}, status=status.HTTP_404_NOT_FOUND)

        return Response({"detail": f"Unknown entity type: {entity_type}"}, status=status.HTTP_400_BAD_REQUEST)


class ArchivePurgeView(APIView):
    """
    Permanently hard-deletes an archived entity. Gated strictly to ADMIN.
    """
    permission_classes = [IsAdminOnly]

    def delete(self, request, entity_type, pk):
        from accounts.audit import log_audit_event
        entity_type = entity_type.upper()

        if entity_type in ["VACANCY", "JOB"]:
            try:
                item = Job.objects.get(pk=pk, is_archived=True)
                title = item.title
                item.delete()
                log_audit_event(
                    request,
                    "RECORD_PURGED",
                    f"Permanently purged archived vacancy '{title}' (ID: {pk})",
                    metadata={"entity_type": "VACANCY", "id": pk, "title": title}
                )
                return Response({"detail": f"Vacancy '{title}' has been permanently purged."})
            except Job.DoesNotExist:
                return Response({"detail": "Archived vacancy not found or is not archived."}, status=status.HTTP_404_NOT_FOUND)

        elif entity_type == "APPLICATION":
            try:
                item = Application.objects.get(pk=pk, is_archived=True)
                desc = f"Application #{pk} for {item.user.username}"
                # A completed attachment is reachable through this URL too, and
                # purging it here would bypass the ATTACHMENT purge guard while
                # destroying the same history, so it is refused identically.
                if is_completed_attachment(item):
                    return Response(
                        {
                            "detail": (
                                f"{desc} is a completed attachment and is permanent institutional "
                                "history; it cannot be purged."
                            )
                        },
                        status=status.HTTP_400_BAD_REQUEST,
                    )
                item.delete()
                log_audit_event(
                    request,
                    "RECORD_PURGED",
                    f"Permanently purged archived {desc}",
                    metadata={"entity_type": "APPLICATION", "id": pk}
                )
                return Response({"detail": f"{desc} has been permanently purged."})
            except Application.DoesNotExist:
                return Response({"detail": "Archived application not found or is not archived."}, status=status.HTTP_404_NOT_FOUND)

        elif entity_type == "ATTACHMENT":
            # Refused on purpose. Purging the underlying application cascades to
            # the deployment, the clearance record and the attachee's uploaded
            # final report, which is exactly the history the archive was created
            # to keep referable. An attachee record can be restored, but it is
            # never destroyed.
            return Response(
                {
                    "detail": (
                        "Completed attachee records are permanent institutional history and cannot be "
                        "purged. Use restore if this record was archived in error."
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        elif entity_type == "DEPARTMENT":
            try:
                item = Department.objects.get(pk=pk, is_archived=True)
                name = item.name
                item.delete()
                log_audit_event(
                    request,
                    "RECORD_PURGED",
                    f"Permanently purged archived department '{name}' (ID: {pk})",
                    metadata={"entity_type": "DEPARTMENT", "id": pk, "name": name}
                )
                return Response({"detail": f"Department '{name}' has been permanently purged."})
            except Department.DoesNotExist:
                return Response({"detail": "Archived department not found or is not archived."}, status=status.HTTP_404_NOT_FOUND)

        return Response({"detail": f"Unknown entity type: {entity_type}"}, status=status.HTTP_400_BAD_REQUEST)


